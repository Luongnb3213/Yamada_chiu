from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any


ROOT_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))


DEFAULT_APP_ID = "jp.co.unisys.yamadamobile"
DEFAULT_HOST = "com.opa334.CraneApplication"
DEFAULT_FRIDA_PYTHON = "/Users/macbook/Library/Application Support/pipx/venvs/frida-tools/bin/python"

CRANE_COLUMNS = [
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "crane_assigned_at",
    "crane_last_used_at",
    "frida_device_id",
    "frida_device_name",
]

CRANE_JS = r"""
function sym(n) {
  return Module.getGlobalExportByName ? Module.getGlobalExportByName(n) : Module.getExportByName(null, n);
}

const objc_lookUpClass = new NativeFunction(sym("objc_lookUpClass"), "pointer", ["pointer"]);
const sel_registerName = new NativeFunction(sym("sel_registerName"), "pointer", ["pointer"]);
const dlopen = new NativeFunction(sym("dlopen"), "pointer", ["pointer", "int"]);
const access = new NativeFunction(sym("access"), "int", ["pointer", "int"]);
const opendir = new NativeFunction(sym("opendir"), "pointer", ["pointer"]);
const readdir = new NativeFunction(sym("readdir"), "pointer", ["pointer"]);
const closedir = new NativeFunction(sym("closedir"), "int", ["pointer"]);
const msg0 = new NativeFunction(sym("objc_msgSend"), "pointer", ["pointer", "pointer"]);
const msg1 = new NativeFunction(sym("objc_msgSend"), "pointer", ["pointer", "pointer", "pointer"]);
const msg2 = new NativeFunction(sym("objc_msgSend"), "pointer", ["pointer", "pointer", "pointer", "pointer"]);
const msgDisplay = new NativeFunction(sym("objc_msgSend"), "pointer", ["pointer", "pointer", "pointer", "pointer", "bool"]);
const msgVoid1 = new NativeFunction(sym("objc_msgSend"), "void", ["pointer", "pointer", "pointer"]);
const msgVoid2 = new NativeFunction(sym("objc_msgSend"), "void", ["pointer", "pointer", "pointer", "pointer"]);
const msgBool1 = new NativeFunction(sym("objc_msgSend"), "bool", ["pointer", "pointer", "pointer"]);
const msgCount = new NativeFunction(sym("objc_msgSend"), "ulong", ["pointer", "pointer"]);
const msgAt = new NativeFunction(sym("objc_msgSend"), "pointer", ["pointer", "pointer", "ulong"]);

function SEL(s) { return sel_registerName(Memory.allocUtf8String(s)); }
function cls(name) { return objc_lookUpClass(Memory.allocUtf8String(name)); }

const S_shared = SEL("sharedManager");
const S_stringWith = SEL("stringWithUTF8String:");
const S_utf8 = SEL("UTF8String");
const S_count = SEL("count");
const S_objectAtIndex = SEL("objectAtIndex:");
const S_isSupported = SEL("isApplicationSupportedByCrane:");
const S_containerIds = SEL("containerIdentifiersOfApplicationWithIdentifier:");
const S_activeContainer = SEL("activeContainerIdentifierForApplicationWithIdentifier:");
const S_setActiveContainer = SEL("setActiveContainerIdentifier:forApplicationWithIdentifier:");
const S_createContainer = SEL("createNewContainerWithName:forApplicationWithIdentifier:");
const S_deleteContainer = SEL("deleteContainerWithIdentifier:forApplicationWithIdentifier:");
const S_deleteContent = SEL("deleteContentOfContainerWithIdentifier:forApplicationWithIdentifier:");
const S_wipeContainer = SEL("wipeContainerWithIdentifier:forApplicationWithIdentifier:shouldRepopulate:");
const S_reloadApp = SEL("reloadApplicationWithIdentifier:");
const S_displayName = SEL("displayNameForContainerWithIdentifier:ofApplicationWithIdentifier:shouldUseShortVersion:");

let NSString = cls("NSString");

function nsstr(value) {
  if (NSString.isNull()) NSString = cls("NSString");
  if (NSString.isNull()) throw new Error("NSString class not found");
  return msg1(NSString, S_stringWith, Memory.allocUtf8String(String(value || "")));
}

function toStr(ns) {
  if (!ns || ns.isNull()) return "";
  const c = msg0(ns, S_utf8);
  return c.isNull() ? "" : c.readUtf8String();
}

function fileExists(path) {
  try {
    return access(Memory.allocUtf8String(path), 0) === 0;
  } catch (_) {
    return false;
  }
}

function listDir(path) {
  const items = [];
  let dir = ptr(0);
  try {
    dir = opendir(Memory.allocUtf8String(path));
    if (dir.isNull()) return items;
    while (true) {
      const ent = readdir(dir);
      if (ent.isNull()) break;
      // Darwin dirent: d_name starts after ino/seekoff/reclen/namlen/type.
      const name = ent.add(21).readUtf8String();
      if (name && name !== "." && name !== "..") items.push(name);
      if (items.length > 4096) break;
    }
  } catch (_) {
    return items;
  } finally {
    if (dir && !dir.isNull()) {
      try { closedir(dir); } catch (_) {}
    }
  }
  return items;
}

function cranePathCandidates() {
  const seen = {};
  const paths = [];
  function add(path) {
    if (!path || seen[path]) return;
    seen[path] = true;
    paths.push(path);
  }
  function addLibSet(root) {
    add(root + "/usr/lib/libcrane.dylib");
    add(root + "/usr/lib/libCrane.dylib");
    add(root + "/procursus/usr/lib/libcrane.dylib");
    add(root + "/procursus/usr/lib/libCrane.dylib");
    add(root + "/Library/MobileSubstrate/DynamicLibraries/Crane.dylib");
    add(root + "/Library/MobileSubstrate/DynamicLibraries/libcrane.dylib");
  }

  [
    "/var/jb",
    "/private/var/jb",
    "/usr",
    "",
    "/private/var/containers/Bundle/Application/.jbroot-92514976D4D84BBB"
  ].forEach(function (root) {
    if (root) addLibSet(root);
  });
  add("/usr/lib/libcrane.dylib");
  add("/usr/lib/libCrane.dylib");
  add("/Library/MobileSubstrate/DynamicLibraries/Crane.dylib");
  add("/var/jb/Library/MobileSubstrate/DynamicLibraries/Crane.dylib");

  ["/private/var/containers/Bundle/Application", "/var/containers/Bundle/Application"].forEach(function (base) {
    listDir(base).forEach(function (entry) {
      if (entry.indexOf(".jbroot") === 0) addLibSet(base + "/" + entry);
    });
  });

  ["/private/preboot", "/var/preboot"].forEach(function (base) {
    listDir(base).forEach(function (entry) {
      const root = base + "/" + entry;
      addLibSet(root);
      listDir(root).forEach(function (child) {
        const lower = String(child).toLowerCase();
        if (lower.indexOf("jb") >= 0 || lower.indexOf("procursus") >= 0 || lower.indexOf("dopamine") >= 0) {
          addLibSet(root + "/" + child);
        }
      });
    });
  });

  return paths.filter(function (path) { return fileExists(path); }).concat(paths.filter(function (path) { return !fileExists(path); }));
}

function loadCraneClass() {
  let CM = cls("CraneManager");
  if (!CM.isNull()) return CM;

  const paths = cranePathCandidates();

  for (const path of paths) {
    try {
      dlopen(Memory.allocUtf8String(path), 2);
      CM = cls("CraneManager");
      if (!CM.isNull()) return CM;
    } catch (_) {}
  }
  return CM;
}

const CraneManager = loadCraneClass();
const MGR = CraneManager.isNull() ? ptr(0) : msg0(CraneManager, S_shared);

function requireManager() {
  if (MGR.isNull()) throw new Error("CraneManager unavailable in CraneApplication host");
  return MGR;
}

function arrayToStrings(array) {
  if (!array || array.isNull()) return [];
  const total = Number(msgCount(array, S_count));
  const out = [];
  for (let i = 0; i < total; i++) {
    out.push(toStr(msgAt(array, S_objectAtIndex, i)));
  }
  return out;
}

function displayName(bundle, containerId) {
  if (!containerId) return "";
  try {
    return toStr(msgDisplay(requireManager(), S_displayName, nsstr(containerId), nsstr(bundle), true));
  } catch (_) {
    return "";
  }
}

function profileOrdinal(bundle, containerId) {
  if (!containerId || containerId === "DEFAULT") return 0;
  try {
    const containers = arrayToStrings(msg1(requireManager(), S_containerIds, nsstr(bundle)))
      .filter(function (id) { return id !== "DEFAULT"; });
    const index = containers.indexOf(containerId);
    return index >= 0 ? index + 1 : 0;
  } catch (_) {
    return 0;
  }
}

function containerLabel(bundle, containerId) {
  const shown = displayName(bundle, containerId);
  if (shown) return shown;
  const ordinal = profileOrdinal(bundle, containerId);
  return ordinal ? "Profile " + ordinal : "";
}

function info(bundle) {
  const mgr = requireManager();
  const app = String(bundle || "");
  const containers = arrayToStrings(msg1(mgr, S_containerIds, nsstr(app)));
  const active = toStr(msg1(mgr, S_activeContainer, nsstr(app)));
  const activeName = displayName(app, active);
  let supported = true;
  try {
    supported = msgBool1(mgr, S_isSupported, nsstr(app));
  } catch (_) {}
  return {
    ok: true,
    app_id: app,
    supported: supported,
    active_container_id: active,
    active_container_name: activeName,
    active_container_label: activeName || containerLabel(app, active),
    containers: containers,
    count: containers.length
  };
}

function create(bundle, name) {
  const app = String(bundle || "");
  const containerName = String(name || "");
  const id = toStr(msg2(requireManager(), S_createContainer, nsstr(containerName), nsstr(app)));
  return {
    ok: !!id,
    app_id: app,
    crane_container_id: id,
    crane_container_name: containerName,
    active_container_id: toStr(msg1(requireManager(), S_activeContainer, nsstr(app)))
  };
}

function switchTo(bundle, containerId, reload) {
  const app = String(bundle || "");
  const id = String(containerId || "").trim();
  if (!id) throw new Error("Missing container id");
  msgVoid2(requireManager(), S_setActiveContainer, nsstr(id), nsstr(app));
  if (reload) msgVoid1(requireManager(), S_reloadApp, nsstr(app));
  const shownName = displayName(app, id);
  return {
    ok: true,
    app_id: app,
    crane_container_id: id,
    crane_container_name: shownName,
    crane_container_label: shownName || containerLabel(app, id),
    active_container_id: toStr(msg1(requireManager(), S_activeContainer, nsstr(app))),
    reloaded: !!reload
  };
}

function ensure(bundle, name, existingId, reload) {
  const app = String(bundle || "");
  const id = String(existingId || "").trim();
  if (id) return switchTo(app, id, reload);

  const created = create(app, name);
  if (!created.ok) return created;
  const switched = switchTo(app, created.crane_container_id, reload);
  switched.crane_container_name = created.crane_container_name;
  switched.created = true;
  return switched;
}

function nextActive(bundle, name, reload, skipDefault, wrap, createIfExhausted) {
  const app = String(bundle || "");
  const before = info(app);
  const containers = before.containers.filter(function (id) {
    return !(skipDefault && id === "DEFAULT");
  });
  if (!containers.length) {
    if (!createIfExhausted) return { ok: false, app_id: app, error: "No containers available" };
    const created = create(app, name);
    if (!created.ok) return created;
    const switched = switchTo(app, created.crane_container_id, reload);
    switched.crane_container_name = created.crane_container_name;
    switched.previous_container_id = before.active_container_id;
    switched.created = true;
    switched.selected_by = "next-active";
    return switched;
  }

  let activeIndex = containers.indexOf(before.active_container_id);
  let nextIndex = activeIndex < 0 ? 0 : activeIndex + 1;
  if (nextIndex >= containers.length) {
    if (wrap) {
      nextIndex = 0;
    } else if (createIfExhausted) {
      const created = create(app, name);
      if (!created.ok) return created;
      const switched = switchTo(app, created.crane_container_id, reload);
      switched.crane_container_name = created.crane_container_name;
      switched.previous_container_id = before.active_container_id;
      switched.created = true;
      switched.selected_by = "next-active";
      switched.exhausted_existing = true;
      return switched;
    } else {
      return {
        ok: false,
        app_id: app,
        error: "Active container is the last container; pass --wrap-containers or --create-if-exhausted.",
        active_container_id: before.active_container_id,
        count: containers.length
      };
    }
  }

  const target = containers[nextIndex];
  const switched = switchTo(app, target, reload);
  switched.previous_container_id = before.active_container_id;
  switched.selected_by = "next-active";
  switched.container_index = nextIndex;
  switched.container_count = containers.length;
  switched.skipped_default = !!skipDefault;
  return switched;
}

function deleteContainer(bundle, containerId) {
  const app = String(bundle || "");
  const id = String(containerId || "").trim();
  if (!id) throw new Error("Missing container id");
  msgVoid2(requireManager(), S_deleteContainer, nsstr(id), nsstr(app));
  return {
    ok: true,
    app_id: app,
    crane_container_id: id,
    deleted: true,
    containers: info(app).containers
  };
}

function deleteContent(bundle, containerId) {
  const app = String(bundle || "");
  const id = String(containerId || "").trim();
  if (!id) throw new Error("Missing container id");
  msgVoid2(requireManager(), S_deleteContent, nsstr(id), nsstr(app));
  return {
    ok: true,
    app_id: app,
    crane_container_id: id,
    content_deleted: true
  };
}

function wipeContainer(bundle, containerId, repopulate) {
  const app = String(bundle || "");
  const id = String(containerId || "").trim();
  if (!id) throw new Error("Missing container id");
  const msgVoid2Bool = new NativeFunction(sym("objc_msgSend"), "void", ["pointer", "pointer", "pointer", "pointer", "bool"]);
  msgVoid2Bool(requireManager(), S_wipeContainer, nsstr(id), nsstr(app), !!repopulate);
  return {
    ok: true,
    app_id: app,
    crane_container_id: id,
    wiped: true,
    repopulate: !!repopulate
  };
}

rpc.exports = {
  ready: function () { return !MGR.isNull(); },
  info: function (bundle) {
    try { return info(bundle); } catch (e) { return { ok: false, error: String(e) }; }
  },
  list: function (bundle) {
    try { return info(bundle).containers; } catch (e) { return { ok: false, error: String(e) }; }
  },
  active: function (bundle) {
    try {
      const data = info(bundle);
      return {
        ok: true,
        app_id: data.app_id,
        active_container_id: data.active_container_id,
        active_container_name: data.active_container_name,
        active_container_label: data.active_container_label
      };
    } catch (e) { return { ok: false, error: String(e) }; }
  },
  create: function (bundle, name) {
    try { return create(bundle, name); } catch (e) { return { ok: false, error: String(e) }; }
  },
  switchto: function (bundle, containerId, reload) {
    try { return switchTo(bundle, containerId, !!reload); } catch (e) { return { ok: false, error: String(e) }; }
  },
  ensure: function (bundle, name, existingId, reload) {
    try { return ensure(bundle, name, existingId, !!reload); } catch (e) { return { ok: false, error: String(e) }; }
  },
  nextactive: function (bundle, name, reload, skipDefault, wrap, createIfExhausted) {
    try { return nextActive(bundle, name, !!reload, !!skipDefault, !!wrap, !!createIfExhausted); } catch (e) { return { ok: false, error: String(e) }; }
  },
  deletecontainer: function (bundle, containerId) {
    try { return deleteContainer(bundle, containerId); } catch (e) { return { ok: false, error: String(e) }; }
  },
  deletecontent: function (bundle, containerId) {
    try { return deleteContent(bundle, containerId); } catch (e) { return { ok: false, error: String(e) }; }
  },
  wipecontainer: function (bundle, containerId, repopulate) {
    try { return wipeContainer(bundle, containerId, !!repopulate); } catch (e) { return { ok: false, error: String(e) }; }
  }
};
"""


def json_print(payload: Any) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2))


def import_frida():
    try:
        import frida  # type: ignore
    except ModuleNotFoundError:
        return None
    return frida


def candidate_frida_pythons() -> list[str]:
    candidates = [
        os.environ.get("FRIDA_PYTHON", ""),
        DEFAULT_FRIDA_PYTHON,
        shutil.which("python3") or "",
        sys.executable,
    ]
    seen: set[str] = set()
    out: list[str] = []
    for candidate in candidates:
        if not candidate:
            continue
        path = str(Path(candidate).expanduser())
        if path in seen or not Path(path).exists():
            continue
        seen.add(path)
        out.append(path)
    return out


def can_import_frida(python_bin: str) -> bool:
    cmd = [python_bin, "-c", "import frida"]
    try:
        return subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False).returncode == 0
    except OSError:
        return False


def find_frida_python() -> str:
    for python_bin in candidate_frida_pythons():
        if can_import_frida(python_bin):
            return python_bin
    raise RuntimeError(
        "Không tìm thấy Python có module frida. Có thể set FRIDA_PYTHON trỏ tới venv frida-tools."
    )


def frida_device_summary(frida) -> str:
    try:
        devices = frida.enumerate_devices()
    except Exception as exc:
        return f"không enumerate được device: {exc}"
    if not devices:
        return "không có device nào"
    return ", ".join(f"{device.id}/{device.name}/{device.type}" for device in devices)


def resolve_frida_device(frida, device_id: str, timeout: int = 10):
    if device_id and device_id.lower() != "auto":
        return frida.get_device(device_id)

    last_error = None
    for attempt in range(1, 4):
        try:
            return frida.get_usb_device(timeout=timeout if attempt == 1 else 3)
        except Exception as exc:
            last_error = exc
            if attempt < 3:
                print(f"[crane] Chưa thấy iPhone USB qua Frida, thử lại {attempt}/3...", file=sys.stderr)
                time.sleep(1.5)

    devices = frida_device_summary(frida)
    raise RuntimeError(
        "Frida chưa thấy iPhone USB. Kiểm tra cáp/Trust This Computer/frida-server trên iPhone. "
        f"Devices hiện thấy: {devices}. Lỗi gốc: {last_error}"
    )


class CraneHostSession:
    def __init__(self, device_id: str, host: str, keep_host: bool = False):
        self.device_id = device_id
        self.host = host
        self.keep_host = keep_host
        self.frida = None
        self.device = None
        self.pid: int | None = None
        self.session = None
        self.script = None

    def __enter__(self):
        frida = import_frida()
        if frida is None:
            raise RuntimeError("Current Python cannot import frida.")
        self.frida = frida
        self.device = resolve_frida_device(frida, self.device_id, timeout=10)

        self.pid = self.device.spawn([self.host])
        self.session = self.device.attach(self.pid)
        self.script = self.session.create_script(CRANE_JS)
        self.script.on("message", self._on_message)
        self.script.load()
        return self

    def __exit__(self, exc_type, exc, tb):
        try:
            if self.session is not None:
                self.session.detach()
        finally:
            if not self.keep_host and self.device is not None and self.pid is not None:
                try:
                    self.device.kill(self.pid)
                except Exception:
                    pass

    def _on_message(self, message, data):
        if message.get("type") == "error":
            print(message.get("stack") or message, file=sys.stderr)

    @property
    def exports(self):
        if self.script is None:
            raise RuntimeError("Frida script is not loaded.")
        return self.script.exports_sync

    def call(self, action: str, app_id: str, **kwargs) -> Any:
        ex = self.exports
        if not ex.ready():
            raise RuntimeError("Không có CraneManager trong CraneApplication host.")
        result: Any
        if action == "info":
            result = ex.info(app_id)
        elif action == "list":
            result = ex.list(app_id)
        elif action == "active":
            result = ex.active(app_id)
        elif action == "create":
            result = ex.create(app_id, kwargs.get("name", ""))
        elif action == "switch":
            result = ex.switchto(app_id, kwargs.get("container_id", ""), bool(kwargs.get("reload")))
        elif action == "ensure":
            result = ex.ensure(
                app_id,
                kwargs.get("name", ""),
                kwargs.get("container_id", ""),
                bool(kwargs.get("reload")),
            )
        elif action == "next":
            result = ex.nextactive(
                app_id,
                kwargs.get("name", ""),
                bool(kwargs.get("reload")),
                bool(kwargs.get("skip_default")),
                bool(kwargs.get("wrap")),
                bool(kwargs.get("create_if_exhausted")),
            )
        elif action == "delete":
            result = ex.deletecontainer(app_id, kwargs.get("container_id", ""))
        elif action == "delete-content":
            result = ex.deletecontent(app_id, kwargs.get("container_id", ""))
        elif action == "wipe":
            result = ex.wipecontainer(
                app_id,
                kwargs.get("container_id", ""),
                bool(kwargs.get("repopulate")),
            )
        else:
            raise RuntimeError(f"Unsupported Crane action: {action}")
        if isinstance(result, dict) and self.device is not None:
            result["frida_device_id"] = getattr(self.device, "id", "")
            result["frida_device_name"] = getattr(self.device, "name", "")
        return result


def invoke_crane_rpc(args: argparse.Namespace, action: str, **kwargs) -> Any:
    if import_frida() is not None:
        with CraneHostSession(args.device_id, args.host, args.keep_host) as crane:
            return crane.call(action, args.app_id, **kwargs)

    python_bin = find_frida_python()
    child_args = [
        python_bin,
        str(Path(__file__).resolve()),
        "_rpc",
        "--action",
        action,
        "--app-id",
        args.app_id,
        "--host",
        args.host,
        "--device-id",
        args.device_id,
    ]
    if args.keep_host:
        child_args.append("--keep-host")
    for key, value in kwargs.items():
        if value is None:
            continue
        if isinstance(value, bool):
            if value:
                child_args.append(f"--{key.replace('_', '-')}")
            continue
        child_args.extend([f"--{key.replace('_', '-')}", str(value)])

    completed = subprocess.run(child_args, text=True, capture_output=True, check=False)
    if completed.returncode != 0:
        stderr = completed.stderr.strip()
        stdout = completed.stdout.strip()
        if stderr.startswith("[crane] "):
            stderr = stderr[len("[crane] "):]
        raise RuntimeError(stderr or stdout or f"Crane RPC failed with exit code {completed.returncode}.")
    try:
        return json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Crane RPC did not return JSON: {completed.stdout!r}") from exc


def sanitize_container_name(value: str) -> str:
    cleaned = re.sub(r"@.*$", "", value or "")
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", cleaned).strip("_")
    return cleaned[:48] or "account"


def default_container_name(profile: dict, row_number: int | None = None) -> str:
    base = sanitize_container_name(str(profile.get("email") or ""))
    suffix = f"r{row_number}" if row_number else datetime.now().strftime("%Y%m%d%H%M%S")
    return f"Chiu_{base}_{suffix}"


def normalize_header(value: object) -> str:
    from src.connections.xlsx_connection import COLUMN_ALIASES

    text = str(value or "").strip().lower()
    return COLUMN_ALIASES.get(text, text)


def ensure_crane_headers(ws) -> dict[str, int]:
    from src.connections.xlsx_connection import INPUT_HEADERS

    headers = [normalize_header(cell.value) for cell in ws[1]]
    if not headers or "email" not in headers:
        ws.delete_rows(1, ws.max_row)
        ws.append(INPUT_HEADERS)
        headers = [normalize_header(cell.value) for cell in ws[1]]

    for column in CRANE_COLUMNS:
        if column not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=column)
            headers.append(column)
    return {header: index + 1 for index, header in enumerate(headers) if header}


def write_crane_assignment(
    xlsx_path: Path,
    sheet_name: str,
    row_number: int,
    container_id: str,
    container_name: str,
    status: str,
    device_id: str = "",
    device_name: str = "",
) -> dict:
    import openpyxl
    from src.connections.xlsx_connection import excel_write_lock

    with excel_write_lock(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        try:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            elif "Mail" in wb.sheetnames:
                ws = wb["Mail"]
            else:
                ws = wb[wb.sheetnames[0]]
            col_map = ensure_crane_headers(ws)
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            ws.cell(row=row_number, column=col_map["crane_container_id"], value=container_id)
            ws.cell(row=row_number, column=col_map["crane_container_name"], value=container_name)
            ws.cell(row=row_number, column=col_map["crane_status"], value=status)
            if not ws.cell(row=row_number, column=col_map["crane_assigned_at"]).value:
                ws.cell(row=row_number, column=col_map["crane_assigned_at"], value=now)
            ws.cell(row=row_number, column=col_map["crane_last_used_at"], value=now)
            ws.cell(row=row_number, column=col_map["frida_device_id"], value=device_id)
            ws.cell(row=row_number, column=col_map["frida_device_name"], value=device_name)

            tmp = xlsx_path.with_name(f".{xlsx_path.stem}.{os.getpid()}.{time.time_ns()}.tmp.xlsx")
            try:
                wb.save(tmp)
                tmp.replace(xlsx_path)
            finally:
                if tmp.exists():
                    try:
                        tmp.unlink()
                    except Exception:
                        pass
            return {
                "xlsx": str(xlsx_path),
                "sheet": ws.title,
                "row": row_number,
                "crane_container_id": container_id,
                "crane_container_name": container_name,
                "crane_status": status,
                "frida_device_id": device_id,
                "frida_device_name": device_name,
            }
        finally:
            wb.close()


def assigned_container_rows(xlsx_path: Path, current_sheet: str, current_row: int) -> dict[str, list[dict[str, Any]]]:
    import openpyxl

    assigned: dict[str, list[dict[str, Any]]] = {}
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        for ws in wb.worksheets:
            headers = [normalize_header(cell.value) for cell in ws[1]]
            if "crane_container_id" not in headers:
                continue
            container_pos = headers.index("crane_container_id")
            email_pos = headers.index("email") if "email" in headers else None
            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if ws.title == current_sheet and row_number == current_row:
                    continue
                container_id = str(row[container_pos] if container_pos < len(row) else "" or "").strip()
                if not container_id:
                    continue
                email = str(row[email_pos] if email_pos is not None and email_pos < len(row) else "" or "").strip()
                assigned.setdefault(container_id, []).append({
                    "sheet": ws.title,
                    "row": row_number,
                    "email": email,
                })
    finally:
        wb.close()
    return assigned


def excel_lock_path(xlsx_path: Path) -> Path:
    return xlsx_path.parent / f".~lock.{xlsx_path.name}#"


def ensure_row(args: argparse.Namespace) -> dict:
    from chiu_profile_from_excel import load_row, profile_from_record

    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        raise RuntimeError(f"XLSX not found: {xlsx_path}")
    lock_path = excel_lock_path(xlsx_path)
    if args.write_excel and lock_path.exists():
        raise RuntimeError(f"Excel appears to be open/locked: {lock_path}. Đóng file rồi chạy lại.")

    record, row_number, sheet_name = load_row(xlsx_path, args.sheet, args.row, args.email)
    profile = profile_from_record(record)
    existing_id = (args.container_id or profile.get("crane_container_id") or "").strip()
    if args.container_mode == "create":
        existing_id = ""
        profile.pop("crane_container_name", None)
    row_device_id = str(record.get("frida_device_id") or "").strip()
    requested_device_id = str(args.device_id or "").strip()
    if row_device_id and requested_device_id != row_device_id:
        print(
            f"[crane] Ưu tiên frida_device_id trong Excel: {row_device_id} "
            f"(bỏ qua lựa chọn {requested_device_id or 'auto'}).",
            file=sys.stderr,
        )
        args.device_id = row_device_id
    name = (
        args.name
        or profile.get("crane_container_name")
        or default_container_name(profile, row_number)
    )

    if existing_id:
        result = invoke_crane_rpc(
            args,
            "ensure",
            name=name,
            container_id=existing_id,
            reload=not args.no_reload,
        )
    elif args.container_mode == "active-then-create":
        info = invoke_crane_rpc(args, "info")
        if not isinstance(info, dict) or not info.get("ok"):
            raise RuntimeError(str(info.get("error") if isinstance(info, dict) else info))
        active_id = str(info.get("active_container_id") or "").strip()
        assigned_rows = assigned_container_rows(xlsx_path, sheet_name, row_number)
        if active_id and not assigned_rows.get(active_id):
            result = invoke_crane_rpc(
                args,
                "ensure",
                name=name,
                container_id=active_id,
                reload=not args.no_reload,
            )
            if isinstance(result, dict):
                result["selected_by"] = "active-current"
                result["previous_container_id"] = active_id
        else:
            result = invoke_crane_rpc(args, "ensure", name=name, container_id="", reload=not args.no_reload)
            if isinstance(result, dict):
                result["selected_by"] = "created-after-active-unavailable"
                result["previous_container_id"] = active_id
                result["active_was_assigned_to"] = assigned_rows.get(active_id, []) if active_id else []
    elif args.container_mode == "active-then-next":
        info = invoke_crane_rpc(args, "info")
        if not isinstance(info, dict) or not info.get("ok"):
            raise RuntimeError(str(info.get("error") if isinstance(info, dict) else info))
        active_id = str(info.get("active_container_id") or "").strip()
        containers = [str(item or "").strip() for item in info.get("containers", []) if str(item or "").strip()]
        assigned_rows = assigned_container_rows(xlsx_path, sheet_name, row_number)
        selected_id = ""
        selected_by = ""

        if active_id and not assigned_rows.get(active_id):
            selected_id = active_id
            selected_by = "active-current"
        else:
            available = [
                container_id for container_id in containers
                if (args.include_default or container_id != "DEFAULT") and not assigned_rows.get(container_id)
            ]
            if active_id in available:
                start = available.index(active_id) + 1
                ordered = available[start:] + available[:start]
            else:
                filtered = [container_id for container_id in containers if args.include_default or container_id != "DEFAULT"]
                try:
                    active_index = filtered.index(active_id)
                except ValueError:
                    active_index = -1
                after_active = filtered[active_index + 1:] if active_index >= 0 else filtered
                before_active = filtered[:active_index + 1] if active_index >= 0 else []
                ordered = [container_id for container_id in after_active + before_active if not assigned_rows.get(container_id)]
            if ordered:
                selected_id = ordered[0]
                selected_by = "next-unused"

        if selected_id:
            result = invoke_crane_rpc(
                args,
                "ensure",
                name=name,
                container_id=selected_id,
                reload=not args.no_reload,
            )
            if isinstance(result, dict):
                result["selected_by"] = selected_by
                result["previous_container_id"] = active_id
                result["active_was_assigned_to"] = assigned_rows.get(active_id, []) if active_id else []
        else:
            if not args.create_if_exhausted:
                raise RuntimeError("Không còn container trống để gán cho row này.")
            result = invoke_crane_rpc(args, "ensure", name=name, container_id="", reload=not args.no_reload)
            if isinstance(result, dict):
                result["selected_by"] = "created-after-exhausted"
                result["previous_container_id"] = active_id
                result["active_was_assigned_to"] = assigned_rows.get(active_id, []) if active_id else []
    elif args.container_mode == "next-active":
        result = invoke_crane_rpc(
            args,
            "next",
            name=name,
            reload=not args.no_reload,
            skip_default=not args.include_default,
            wrap=args.wrap_containers,
            create_if_exhausted=args.create_if_exhausted,
        )
    else:
        result = invoke_crane_rpc(
            args,
            "ensure",
            name=name,
            container_id="",
            reload=not args.no_reload,
        )
    if not isinstance(result, dict):
        raise RuntimeError(f"Unexpected Crane result: {result!r}")
    if not result.get("ok"):
        raise RuntimeError(str(result.get("error") or result))

    container_id = str(result.get("crane_container_id") or result.get("active_container_id") or "")
    container_name = str(result.get("crane_container_name") or result.get("active_container_name") or "")
    container_label = str(
        result.get("crane_container_label")
        or result.get("active_container_label")
        or container_name
        or ""
    )
    if not container_name and result.get("container_index") is not None:
        try:
            container_name = f"Profile {int(result.get('container_index')) + 1}"
        except (TypeError, ValueError):
            container_name = ""
    if not container_name and result.get("created"):
        container_name = name
    if not container_name and container_id == "DEFAULT":
        container_name = "DEFAULT"
    result["crane_container_id"] = container_id
    result["crane_container_name"] = container_name
    result["crane_container_label"] = container_label
    if args.write_excel:
        result["excel"] = write_crane_assignment(
            xlsx_path,
            sheet_name,
            row_number,
            container_id,
            container_name,
            args.status,
            str(result.get("frida_device_id") or args.device_id or ""),
            str(result.get("frida_device_name") or ""),
        )
    result["source"] = {"xlsx": str(xlsx_path), "sheet": sheet_name, "row": row_number}
    return result


def add_common_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--app-id", default=os.environ.get("YAMADA_APP_ID", DEFAULT_APP_ID))
    parser.add_argument("--host", default=os.environ.get("CRANE_HOST", DEFAULT_HOST))
    parser.add_argument("--device-id", default=os.environ.get("FRIDA_DEVICE_ID", "auto"))
    parser.add_argument("--keep-host", action="store_true", help="Do not kill the frozen Crane host process on exit.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Manage Crane containers for Yamada_chiu rows.")
    sub = parser.add_subparsers(dest="command", required=True)

    for name in ("info", "list", "active"):
        add_common_args(sub.add_parser(name))

    create = sub.add_parser("create")
    add_common_args(create)
    create.add_argument("--name", default="")

    switch = sub.add_parser("switch")
    add_common_args(switch)
    switch.add_argument("--container-id", required=True)
    switch.add_argument("--no-reload", action="store_true")

    ensure = sub.add_parser("ensure")
    add_common_args(ensure)
    ensure.add_argument("--container-id", default="")
    ensure.add_argument("--name", default="")
    ensure.add_argument("--no-reload", action="store_true")

    next_container = sub.add_parser("next")
    add_common_args(next_container)
    next_container.add_argument("--name", default="")
    next_container.add_argument("--no-reload", action="store_true")
    next_container.add_argument("--include-default", action="store_true")
    next_container.add_argument("--wrap-containers", action="store_true")
    next_container.add_argument("--create-if-exhausted", action="store_true")

    ensure_excel = sub.add_parser("ensure-row")
    add_common_args(ensure_excel)
    ensure_excel.add_argument("--xlsx", required=True)
    ensure_excel.add_argument("--sheet", default="Accounts")
    ensure_excel.add_argument("--row", type=int)
    ensure_excel.add_argument("--email", default="")
    ensure_excel.add_argument("--container-id", default="")
    ensure_excel.add_argument("--name", default="")
    ensure_excel.add_argument(
        "--container-mode",
        choices=["active-then-create", "active-then-next", "next-active", "create"],
        default="active-then-next",
    )
    ensure_excel.add_argument("--include-default", action="store_true")
    ensure_excel.add_argument("--wrap-containers", action="store_true")
    ensure_excel.add_argument("--no-create-if-exhausted", dest="create_if_exhausted", action="store_false")
    ensure_excel.add_argument("--status", default="ASSIGNED")
    ensure_excel.add_argument("--no-reload", action="store_true")
    ensure_excel.add_argument("--no-write-excel", dest="write_excel", action="store_false")
    ensure_excel.set_defaults(write_excel=True, create_if_exhausted=True)

    delete = sub.add_parser("delete")
    add_common_args(delete)
    delete.add_argument("--container-id", required=True)

    delete_content = sub.add_parser("delete-content")
    add_common_args(delete_content)
    delete_content.add_argument("--container-id", required=True)

    wipe = sub.add_parser("wipe")
    add_common_args(wipe)
    wipe.add_argument("--container-id", required=True)
    wipe.add_argument("--repopulate", action="store_true")

    rpc = sub.add_parser("_rpc")
    add_common_args(rpc)
    rpc.add_argument(
        "--action",
        required=True,
        choices=["info", "list", "active", "create", "switch", "ensure", "next", "delete", "delete-content", "wipe"],
    )
    rpc.add_argument("--name", default="")
    rpc.add_argument("--container-id", default="")
    rpc.add_argument("--reload", action="store_true")
    rpc.add_argument("--repopulate", action="store_true")
    rpc.add_argument("--skip-default", action="store_true")
    rpc.add_argument("--wrap", action="store_true")
    rpc.add_argument("--create-if-exhausted", action="store_true")

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        if args.command == "_rpc":
            result = invoke_crane_rpc(
                args,
                args.action,
                name=getattr(args, "name", ""),
                container_id=getattr(args, "container_id", ""),
                reload=getattr(args, "reload", False),
                repopulate=getattr(args, "repopulate", False),
                skip_default=getattr(args, "skip_default", False),
                wrap=getattr(args, "wrap", False),
                create_if_exhausted=getattr(args, "create_if_exhausted", False),
            )
        elif args.command == "info":
            result = invoke_crane_rpc(args, "info")
        elif args.command == "list":
            result = invoke_crane_rpc(args, "list")
        elif args.command == "active":
            result = invoke_crane_rpc(args, "active")
        elif args.command == "create":
            name = args.name or default_container_name({}, None)
            result = invoke_crane_rpc(args, "create", name=name)
        elif args.command == "switch":
            result = invoke_crane_rpc(args, "switch", container_id=args.container_id, reload=not args.no_reload)
        elif args.command == "ensure":
            name = args.name or default_container_name({}, None)
            result = invoke_crane_rpc(
                args,
                "ensure",
                name=name,
                container_id=args.container_id,
                reload=not args.no_reload,
            )
        elif args.command == "next":
            name = args.name or default_container_name({}, None)
            result = invoke_crane_rpc(
                args,
                "next",
                name=name,
                reload=not args.no_reload,
                skip_default=not args.include_default,
                wrap=args.wrap_containers,
                create_if_exhausted=args.create_if_exhausted,
            )
        elif args.command == "ensure-row":
            result = ensure_row(args)
        elif args.command == "delete":
            result = invoke_crane_rpc(args, "delete", container_id=args.container_id)
        elif args.command == "delete-content":
            result = invoke_crane_rpc(args, "delete-content", container_id=args.container_id)
        elif args.command == "wipe":
            result = invoke_crane_rpc(args, "wipe", container_id=args.container_id, repopulate=args.repopulate)
        else:
            parser.error(f"Unsupported command: {args.command}")
            return 2
    except Exception as exc:
        print(f"[crane] {exc}", file=sys.stderr)
        return 1

    json_print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
