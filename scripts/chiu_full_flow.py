from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import excel_write_lock  # noqa: E402
from chiu_profile_from_excel import load_row, profile_from_record, write_profile_js  # noqa: E402
from card_rotation import apply_card_rotation  # noqa: E402


EMAIL_OTP_PROCESS_TIMEOUT_SECONDS = 120


class FlowResultError(RuntimeError):
    def __init__(self, message: str, *, gold_status: str = "", chiu_status: str = "", notes: str = "") -> None:
        super().__init__(message)
        self.gold_status = gold_status
        self.chiu_status = chiu_status
        self.notes = notes


def quote_cmd(cmd: list[str]) -> str:
    return " ".join(shlex.quote(str(part)) for part in cmd)


def run_cmd_output(cmd: list[str], title: str, stream: bool = False, timeout: int | None = None) -> str:
    print(f"\n--- {title} ---", flush=True)
    print("$ " + quote_cmd(cmd), flush=True)
    proc = subprocess.Popen(
        cmd,
        cwd=str(ROOT_DIR),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
    )
    assert proc.stdout is not None
    try:
        output, _ = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        proc.kill()
        output, _ = proc.communicate()
        if output.strip():
            print(output.strip(), flush=True)
        raise RuntimeError(f"{title} quá timeout {timeout}s.") from exc

    lines = output.splitlines(keepends=True)
    for line in lines:
        if stream:
            print(line, end="", flush=True)
    code = proc.returncode
    if code != 0:
        if not stream and output.strip():
            print(output.strip(), flush=True)
        raise RuntimeError(f"{title} lỗi exit={code}.")
    return output


def print_prefixed_lines(output: str, prefixes: tuple[str, ...]) -> None:
    for line in output.splitlines():
        if line.startswith(prefixes):
            print(line, flush=True)


def parse_json_from_output(output: str) -> Any:
    decoder = json.JSONDecoder()
    candidates: list[tuple[int, Any]] = []
    for index, char in enumerate(output):
        if char != "{":
            continue
        try:
            value, end = decoder.raw_decode(output[index:])
        except json.JSONDecodeError:
            continue
        candidates.append((end, value))
    if not candidates:
        return {}
    return max(candidates, key=lambda item: item[0])[1]


def latest_state(dom_result: Any) -> str:
    if not isinstance(dom_result, dict):
        return ""
    history = dom_result.get("history")
    if isinstance(history, list) and history:
        for item in reversed(history):
            if isinstance(item, dict) and item.get("state"):
                return str(item.get("state") or "")
    return str(dom_result.get("state") or "")


def dom_history(dom_result: Any) -> list[dict[str, Any]]:
    if isinstance(dom_result, dict) and isinstance(dom_result.get("history"), list):
        return [item for item in dom_result["history"] if isinstance(item, dict)]
    if isinstance(dom_result, dict):
        return [dom_result]
    return []


def print_dom_summary(dom_result: Any, title: str) -> None:
    history = dom_history(dom_result)
    if not history:
        print(f"[dom] {title}: không có history.", flush=True)
        return
    print(f"[dom] {title}: {len(history)} bước", flush=True)
    for index, step in enumerate(history, start=1):
        state = step.get("state") or "-"
        action = step.get("action") or "-"
        extra = ""
        if step.get("needs_otp") or action == "need_login_otp":
            extra = " | cần OTP"
        elif action == "fill_onepiece_lottery_form_and_confirm":
            store = step.get("store") if isinstance(step.get("store"), dict) else {}
            shop = store.get("shopName") or store.get("area") or ""
            extra = f" | đã fill form{': ' + shop if shop else ''}"
        elif action == "onepiece_lottery_already_applied":
            extra = " | đã nộp trước đó"
        elif action == "fail_no_retry":
            reason = str(step.get("reason") or "").strip()
            extra = " | FAIL_NO_RETRY" + (f": {reason}" if reason else "")
        elif action == "wait_timeout":
            last = step.get("last") if isinstance(step.get("last"), dict) else {}
            extra = f" | timeout ở {last.get('state') or '?'}"
        elif action == "gold_already_registered_go_home":
            evidence = step.get("goldEvidence") if isinstance(step.get("goldEvidence"), list) else []
            extra = " | gold đã đăng ký"
            if evidence:
                extra += ": " + ", ".join(str(item) for item in evidence[:3])
        elif state == "unknown" or action == "no_action":
            title_value = str(step.get("title") or "").strip()
            url_value = str(step.get("url") or "").strip()
            hint = title_value or url_value
            if hint:
                extra = f" | {hint[:80]}"
        print(f"[dom] {index}. {state} -> {action}{extra}", flush=True)


def registration_completed(dom_result: Any) -> bool:
    """True if this run drove the nick through new-member registration (reg006)."""
    for item in dom_history(dom_result):
        action = str(item.get("action") or "")
        state = str(item.get("state") or "")
        if action == "launch_app_after_registration" or state == "member_register_complete":
            return True
        if item.get("registered") is True:
            return True
    return False


def needs_fresh_otp(dom_result: Any) -> bool:
    if not isinstance(dom_result, dict):
        return False
    items = dom_result.get("history") if isinstance(dom_result.get("history"), list) else [dom_result]
    for item in items:
        if not isinstance(item, dict):
            continue
        if item.get("needs_otp") or item.get("action") == "need_login_otp":
            return True
    return False


def terminal_dom_error(dom_result: Any) -> str:
    if not isinstance(dom_result, dict):
        return "DOM không trả JSON hợp lệ."
    if dom_result.get("ok") is False:
        return str(dom_result.get("reason") or dom_result.get("error") or dom_result.get("action") or "DOM failed")
    history = dom_result.get("history")
    if isinstance(history, list) and history:
        last = next((item for item in reversed(history) if isinstance(item, dict)), {})
        if last and last.get("ok") is False:
            if last.get("needs_otp"):
                return ""
            if last.get("state") == "no_webview_or_no_result":
                return "Chưa thấy WKWebView của app Yamada. App có thể chưa mở xong, đang trắng màn, hoặc Frida attach nhầm/attach quá sớm."
            return str(last.get("reason") or last.get("error") or last.get("action") or "DOM step failed")
    return ""


def fetch_login_otp(xlsx: Path, sheet_name: str, row: int) -> str:
    output = run_cmd_output(
        [
            sys.executable,
            "scripts/fetch_chiu_email_otp.py",
            "--xlsx",
            str(xlsx),
            "--sheet",
            sheet_name,
            "--row",
            str(row),
        ],
        "Lấy OTP login từ email",
        timeout=EMAIL_OTP_PROCESS_TIMEOUT_SECONDS,
    )
    result = parse_json_from_output(output)
    if not isinstance(result, dict):
        return ""
    return str(result.get("otp") or "").strip()


def current_gold_status(xlsx: Path, sheet_name: str, row: int) -> str:
    try:
        record, _, _ = load_row(xlsx, sheet_name, row)
    except Exception:
        return ""
    return str(record.get("gold_status") or "").strip().upper()


def excel_preferred_device_id(xlsx: Path, sheet_name: str, row: int) -> str:
    try:
        record, _, _ = load_row(xlsx, sheet_name, row)
    except Exception:
        return ""
    return str(record.get("frida_device_id") or "").strip()


def is_auto_device_selection(value: str) -> bool:
    raw = (value or "auto").strip().lower()
    return raw in ("", "auto", "all", "*", "tat-ca", "tất-cả")


def ensure_status_headers(ws) -> dict[str, int]:
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    if not any(headers):
        headers = ["email", "status", "error_details"]
        ws.append(headers)
    for header in (
        "crane_container_id",
        "crane_container_name",
        "crane_status",
        "crane_assigned_at",
        "crane_last_used_at",
        "frida_device_id",
        "frida_device_name",
        "reg_status",
        "gold_status",
        "chiu_status",
        "status",
        "error_details",
        "notes",
    ):
        if header not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=header)
            headers.append(header)
    return {header: index + 1 for index, header in enumerate(headers) if header}


def atomic_save_workbook(wb, xlsx: Path) -> None:
    tmp = xlsx.with_name(f".{xlsx.stem}.{os.getpid()}.{time.time_ns()}.tmp.xlsx")
    try:
        wb.save(tmp)
        tmp.replace(xlsx)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass


def append_run_event(args: argparse.Namespace, stage: str, **payload: Any) -> None:
    event_path = (
        Path(str(args.event_log)).expanduser()
        if getattr(args, "event_log", "")
        else ROOT_DIR / "agents" / "runtime" / "chiu_row_progress.jsonl"
    )
    try:
        event_path.parent.mkdir(parents=True, exist_ok=True)
        event = {
            "ts": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "stage": stage,
            "xlsx": str(Path(args.xlsx).expanduser()),
            "sheet": args.sheet,
            "row": args.row,
            "device_id": args.device_id,
            **payload,
        }
        with event_path.open("a") as handle:
            handle.write(json.dumps(event, ensure_ascii=False) + "\n")
    except Exception:
        pass


def write_row_result(
    xlsx: Path,
    sheet_name: str,
    row: int,
    status: str,
    error_details: str = "",
    *,
    reg_status: str = "",
    gold_status: str = "",
    chiu_status: str = "",
    notes: str = "",
    crane_result: dict | None = None,
) -> None:
    with excel_write_lock(xlsx):
        lock_path = xlsx.parent / f".~lock.{xlsx.name}#"
        if lock_path.exists():
            print(f"[excel] Không ghi được status vì Excel đang mở/lock: {lock_path}", flush=True)
            return
        wb = openpyxl.load_workbook(xlsx)
        try:
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
            col = ensure_status_headers(ws)
            crane = crane_result or {}
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            values = {
                "status": status,
                "error_details": error_details,
                "reg_status": reg_status,
                "gold_status": gold_status,
                "chiu_status": chiu_status,
                "notes": notes,
                "crane_container_id": crane.get("crane_container_id") or crane.get("active_container_id") or "",
                "crane_container_name": crane.get("crane_container_name") or crane.get("active_container_name") or "",
                "crane_status": "ASSIGNED" if crane else "",
                "crane_assigned_at": now if crane else "",
                "crane_last_used_at": now if crane else "",
                "frida_device_id": crane.get("frida_device_id") or "",
                "frida_device_name": crane.get("frida_device_name") or "",
            }
            for key, value in values.items():
                if key in ("reg_status", "gold_status", "chiu_status", "notes") and value in (None, ""):
                    continue
                if key.startswith("crane_") or key.startswith("frida_"):
                    if value in (None, ""):
                        continue
                    if key == "crane_assigned_at" and ws.cell(row=row, column=col[key]).value:
                        continue
                ws.cell(row=row, column=col[key], value=value)
            atomic_save_workbook(wb, xlsx)
        finally:
            wb.close()


def build_dom_cmd(args: argparse.Namespace) -> list[str]:
    cmd = [
        sys.executable,
        "scripts/chiu_dom_runner.py",
        "--action",
        "run",
        "--wait-timeout-ms",
        str(args.wait_timeout_ms),
        "--max-steps",
        str(args.max_steps),
        "--flow-complete-delay-ms",
        str(args.flow_complete_delay_ms),
        "--gold-payment-wait-timeout-ms",
        str(args.gold_payment_wait_timeout_ms),
        "--device-id",
        args.device_id,
        "--profile-js",
        str(args.profile_js),
    ]
    if args.no_submit:
        cmd.append("--no-submit")
    if getattr(args, "enable_register", False):
        cmd.append("--enable-register")
    return cmd


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run one Yamada_chiu row through the current membership/chiu entry flow.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Accounts")
    parser.add_argument("--row", type=int, required=True)
    parser.add_argument("--no-reload", action="store_true")
    parser.add_argument("--no-submit", action="store_true")
    parser.add_argument(
        "--enable-register",
        action="store_true",
        help="Bật luồng đăng ký mới (reg001->reg006) tuỳ reg_status. Không bật thì chỉ login/gold/onepiece như cũ.",
    )
    parser.add_argument("--device-id", default=os.environ.get("FRIDA_DEVICE_ID", "auto"))
    parser.add_argument(
        "--container-mode",
        choices=["active-then-create", "active-then-next", "next-active", "create"],
        default="active-then-create",
    )
    parser.add_argument("--profile-js", default="")
    parser.add_argument("--wait-timeout-ms", type=int, default=25000)
    parser.add_argument("--max-steps", type=int, default=40)
    parser.add_argument("--flow-complete-delay-ms", type=int, default=2000)
    parser.add_argument("--gold-payment-wait-timeout-ms", type=int, default=60000)
    parser.add_argument("--event-log", default="", help="JSONL progress file for deferred/batch recovery.")
    parser.add_argument("--defer-excel-write", action="store_true", help="Write progress to --event-log only; batch merges Excel at the end.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    xlsx = Path(args.xlsx).expanduser()
    row_args = ["--xlsx", str(xlsx), "--sheet", args.sheet, "--row", str(args.row)]
    preferred_device_id = excel_preferred_device_id(xlsx, args.sheet, args.row)
    if preferred_device_id and args.device_id != preferred_device_id and is_auto_device_selection(args.device_id):
        original_device_id = args.device_id
        args.device_id = preferred_device_id
        print(
            f"[excel] Ưu tiên frida_device_id trong Excel: {preferred_device_id} "
            f"(bỏ qua lựa chọn {original_device_id})",
            flush=True,
        )
    elif preferred_device_id and args.device_id != preferred_device_id:
        print(
            f"[excel] Row {args.row} gán cho device {preferred_device_id}, "
            f"không chạy bằng device đã chọn {args.device_id}.",
            file=sys.stderr,
            flush=True,
        )
        return 2
    if args.profile_js:
        profile_js = Path(args.profile_js).expanduser()
    else:
        safe_device = re.sub(r"[^A-Za-z0-9_.-]+", "_", args.device_id or "auto")[:48]
        profile_js = ROOT_DIR / "agents" / "runtime" / f"current_profile_{safe_device}_r{args.row}_{os.getpid()}.js"
    args.profile_js = str(profile_js)
    crane_result: dict = {}
    reg_status_val = ""

    try:
        crane_cmd = [
            sys.executable,
            "scripts/crane_container_manager.py",
            "ensure-row",
            *row_args,
            "--device-id",
            args.device_id,
            "--container-mode",
            args.container_mode,
            "--no-write-excel",
        ]
        if args.no_reload:
            crane_cmd.append("--no-reload")
        crane_output = run_cmd_output(crane_cmd, "Chuẩn bị container")
        crane_result = parse_json_from_output(crane_output)
        if not isinstance(crane_result, dict):
            crane_result = {}
        crane_result.setdefault("frida_device_id", args.device_id)
        append_run_event(args, "crane", crane_result=crane_result)
        print(
            "[crane] container="
            f"{crane_result.get('crane_container_name') or crane_result.get('active_container_name') or crane_result.get('crane_container_label') or crane_result.get('active_container_label') or ''} "
            f"({crane_result.get('crane_container_id') or crane_result.get('active_container_id') or ''})",
            flush=True,
        )

        profile_output = run_cmd_output(
            [
                sys.executable,
                "scripts/chiu_profile_from_excel.py",
                *row_args,
                "--out",
                str(profile_js),
            ],
            "Đọc context từ Excel",
        )
        profile_result = parse_json_from_output(profile_output)
        profile = profile_result.get("profile") if isinstance(profile_result, dict) else {}
        print(f"[excel] email={profile.get('email') or ''} row={args.row}", flush=True)

        dom_output = run_cmd_output(build_dom_cmd(args), "Chạy DOM Gold/Chiu")
        print_prefixed_lines(dom_output, ("[chiu-dom]",))
        dom_result = parse_json_from_output(dom_output)
        print_dom_summary(dom_result, "lượt đầu")
        if isinstance(dom_result, dict) and needs_fresh_otp(dom_result):
            otp = fetch_login_otp(xlsx, args.sheet, args.row)
            if not otp:
                raise RuntimeError("Không lấy được OTP login từ email.")
            record, _, _ = load_row(xlsx, args.sheet, args.row)
            profile_with_otp = profile_from_record(record)
            apply_card_rotation(profile_with_otp, args.row, wait_cooldown=False)  # giữ đúng thẻ xoay theo dòng
            profile_with_otp["auth_code"] = otp
            profile_with_otp["auth_code_source"] = "email_fresh"
            write_profile_js(profile_with_otp, Path(args.profile_js))
            print("[email] Đã lấy OTP login mới từ mail, chạy DOM tiếp.", flush=True)
            dom_output = run_cmd_output(build_dom_cmd(args), "Chạy DOM sau OTP")
            print_prefixed_lines(dom_output, ("[chiu-dom]",))
            dom_result = parse_json_from_output(dom_output)
            print_dom_summary(dom_result, "sau OTP")
        if registration_completed(dom_result):
            reg_status_val = "SUCCESS"
        dom_error = terminal_dom_error(dom_result)
        if dom_error:
            raise RuntimeError(dom_error)

        final_state = latest_state(dom_result)
        if final_state == "chiu_onepiece_submitted":
            append_run_event(
                args,
                "final",
                status="SUCCESS",
                error_details="",
                reg_status=reg_status_val,
                gold_status="SUCCESS",
                chiu_status="SUCCESS",
                notes="gold_done; onepiece_lottery_submitted",
                crane_result=crane_result,
            )
            if not args.defer_excel_write:
                write_row_result(
                    xlsx,
                    args.sheet,
                    args.row,
                    "SUCCESS",
                    "",
                    reg_status=reg_status_val,
                    gold_status="SUCCESS",
                    chiu_status="SUCCESS",
                    notes="gold_done; onepiece_lottery_submitted",
                    crane_result=crane_result,
                )
        elif final_state == "onepiece_lottery_form_filled":
            append_run_event(
                args,
                "final",
                status="FORM_FILLED",
                error_details="",
                reg_status=reg_status_val,
                gold_status="SUCCESS",
                chiu_status="FORM_FILLED",
                notes="gold_done; onepiece_form_filled_no_submit",
                crane_result=crane_result,
            )
            if not args.defer_excel_write:
                write_row_result(
                    xlsx,
                    args.sheet,
                    args.row,
                    "FORM_FILLED",
                    "",
                    reg_status=reg_status_val,
                    gold_status="SUCCESS",
                    chiu_status="FORM_FILLED",
                    notes="gold_done; onepiece_form_filled_no_submit",
                    crane_result=crane_result,
                )
        elif final_state == "ymd_common_error_no_retry":
            append_run_event(
                args,
                "final",
                status="FAIL_NO_RETRY",
                error_details="member_or_system_error",
                reg_status=reg_status_val,
                gold_status="FAIL_NO_RETRY",
                chiu_status="FAIL_NO_RETRY",
                notes="common_error_no_retry",
                crane_result=crane_result,
            )
            if not args.defer_excel_write:
                write_row_result(
                    xlsx,
                    args.sheet,
                    args.row,
                    "FAIL_NO_RETRY",
                    "member_or_system_error",
                    reg_status=reg_status_val,
                    gold_status="FAIL_NO_RETRY",
                    chiu_status="FAIL_NO_RETRY",
                    notes="common_error_no_retry",
                    crane_result=crane_result,
                )
        elif final_state == "gold_card_unusable_no_retry":
            append_run_event(
                args,
                "final",
                status="FAIL_NO_RETRY",
                error_details="card_unusable",
                reg_status=reg_status_val,
                gold_status="FAIL_NO_RETRY",
                chiu_status="FAIL_NO_RETRY",
                notes="card_unusable_no_retry",
                crane_result=crane_result,
            )
            if not args.defer_excel_write:
                write_row_result(
                    xlsx,
                    args.sheet,
                    args.row,
                    "FAIL_NO_RETRY",
                    "card_unusable",
                    reg_status=reg_status_val,
                    gold_status="FAIL_NO_RETRY",
                    chiu_status="FAIL_NO_RETRY",
                    notes="card_unusable_no_retry",
                    crane_result=crane_result,
                )
        elif not args.no_submit and final_state in ("store_sale_tab", "ready_for_chiu_store_sale"):
            raise FlowResultError(
                f"DOM chưa submit One Piece: last_state={final_state}",
                gold_status="SUCCESS",
                chiu_status="FAILED",
                notes="gold_done; chiu_not_finished",
            )
        elif not args.no_submit:
            raise RuntimeError(f"DOM chưa submit One Piece: last_state={final_state}")
        else:
            append_run_event(
                args,
                "final",
                status="SUCCESS",
                error_details="",
                reg_status=reg_status_val,
                notes=f"last_state={final_state}",
                crane_result=crane_result,
            )
            if not args.defer_excel_write:
                write_row_result(xlsx, args.sheet, args.row, "SUCCESS", "", reg_status=reg_status_val, notes=f"last_state={final_state}", crane_result=crane_result)
        print(f"\n[flow] Xong lượt chạy row {args.row} lúc {datetime.now():%Y-%m-%d %H:%M:%S}", flush=True)
        return 0
    except Exception as exc:
        message = str(exc)
        notes = getattr(exc, "notes", "")
        gold_status = getattr(exc, "gold_status", "")
        chiu_status = getattr(exc, "chiu_status", "")
        if not gold_status and not chiu_status and current_gold_status(xlsx, args.sheet, args.row) == "SUCCESS":
            chiu_status = "FAILED"
        elif not gold_status and not chiu_status:
            gold_status = "FAILED"
        append_run_event(
            args,
            "final",
            status="FAILED",
            error_details=message,
            reg_status=reg_status_val,
            gold_status=gold_status,
            chiu_status=chiu_status,
            notes=notes,
            crane_result=crane_result,
        )
        if not args.defer_excel_write:
            write_row_result(
                xlsx,
                args.sheet,
                args.row,
                "FAILED",
                message,
                reg_status=reg_status_val,
                gold_status=gold_status,
                chiu_status=chiu_status,
                notes=notes,
                crane_result=crane_result,
            )
        print(f"[flow] Lỗi: {message}", file=sys.stderr, flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
