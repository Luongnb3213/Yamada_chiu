from __future__ import annotations

import argparse
import json
import os
import queue
import random
import shlex
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_FRIDA_PYTHON = "/Users/macbook/Library/Application Support/pipx/venvs/frida-tools/bin/python"
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import excel_write_lock, normalize_status  # noqa: E402
from chiu_profile_from_excel import cell_text, normalize_header  # noqa: E402


STOP_ON_ERROR_HINTS = (
    "Frida chưa thấy iPhone USB",
    "device not found",
    "unable to find process",
    "Không có CraneManager",
    "CraneManager unavailable",
    "Excel appears to be open/locked",
)


ROW_RESULT_COLUMNS = (
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "crane_assigned_at",
    "crane_last_used_at",
    "frida_device_id",
    "frida_device_name",
    "reg_status",
    "gold_status",
    "chiu_status",
    "status",
    "error_details",
    "notes",
)


def quote_cmd(cmd: list[str]) -> str:
    return " ".join(shlex.quote(str(part)) for part in cmd)


def configure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def child_env() -> dict[str, str]:
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    return env


def choose_sheet(wb, sheet_name: str):
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    for fallback in ("Accounts", "Mail"):
        if fallback in wb.sheetnames:
            return wb[fallback]
    return wb[wb.sheetnames[0]]


def is_auto_device_selection(value: str) -> bool:
    raw = (value or "auto").strip().lower()
    return raw in ("", "auto", "all", "*", "tat-ca", "tất-cả")


def runnable_rows(
    xlsx: Path,
    sheet_name: str,
    limit: int,
    allowed_device_ids: set[str] | None = None,
    balance_device_ids: list[str] | None = None,
) -> tuple[str, list[dict]]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    try:
        ws = choose_sheet(wb, sheet_name)
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return ws.title, []

        headers = [normalize_header(cell) for cell in header_row]
        if "email" not in headers:
            raise RuntimeError(f"Sheet {ws.title!r} không có cột email.")
        email_pos = headers.index("email")
        status_pos = headers.index("status") if "status" in headers else None
        gold_status_pos = headers.index("gold_status") if "gold_status" in headers else None
        chiu_status_pos = headers.index("chiu_status") if "chiu_status" in headers else None
        device_pos = headers.index("frida_device_id") if "frida_device_id" in headers else None
        container_pos = headers.index("crane_container_id") if "crane_container_id" in headers else None

        selected: list[dict] = []
        for row_number, row in enumerate(rows, start=2):
            email = cell_text(row[email_pos]) if email_pos < len(row) else ""
            if not email:
                continue
            status_raw = cell_text(row[status_pos]) if status_pos is not None and status_pos < len(row) else ""
            status = normalize_status(status_raw)
            gold_status_raw = cell_text(row[gold_status_pos]) if gold_status_pos is not None and gold_status_pos < len(row) else status_raw
            chiu_status_raw = cell_text(row[chiu_status_pos]) if chiu_status_pos is not None and chiu_status_pos < len(row) else ""
            gold_status = normalize_status(gold_status_raw)
            chiu_status = normalize_status(chiu_status_raw)
            if "FAIL_NO_RETRY" in (status, gold_status, chiu_status):
                continue
            if chiu_status in ("SUCCESS", "FORM_FILLED") or status in ("SUCCESS", "PROCESSING", "FORM_FILLED"):
                continue
            if status not in ("", "PENDING", "FAILED") and gold_status not in ("", "PENDING", "FAILED"):
                continue
            preferred_device = cell_text(row[device_pos]) if device_pos is not None and device_pos < len(row) else ""
            if allowed_device_ids is not None and preferred_device and preferred_device not in allowed_device_ids:
                continue
            container_id = cell_text(row[container_pos]) if container_pos is not None and container_pos < len(row) else ""
            selected.append({"row": row_number, "device_id": preferred_device, "container_id": container_id})
            if not balance_device_ids and limit > 0 and len(selected) >= limit:
                break
        if balance_device_ids and limit > 0:
            queues = {device_id: [] for device_id in balance_device_ids}
            fallback_index = 0
            for task in selected:
                preferred = str(task.get("device_id") or "")
                if preferred in queues:
                    target = preferred
                else:
                    target = balance_device_ids[fallback_index % len(balance_device_ids)]
                    fallback_index += 1
                queues[target].append(task)
            balanced: list[dict] = []
            while len(balanced) < limit and any(queues.values()):
                for device_id in balance_device_ids:
                    if queues[device_id]:
                        balanced.append(queues[device_id].pop(0))
                        if len(balanced) >= limit:
                            break
            selected = balanced
        elif limit > 0:
            selected = selected[:limit]
        return ws.title, selected
    finally:
        wb.close()


def frida_python() -> str:
    candidates = [
        os.environ.get("FRIDA_PYTHON", ""),
        DEFAULT_FRIDA_PYTHON,
        sys.executable,
        "python3",
    ]
    for candidate in candidates:
        if not candidate:
            continue
        try:
            completed = subprocess.run(
                [candidate, "-c", "import frida"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=False,
                env=child_env(),
            )
            if completed.returncode == 0:
                return candidate
        except OSError:
            continue
    raise RuntimeError("Không tìm thấy Python có module frida để liệt kê device.")


def list_usb_devices() -> list[dict]:
    code = (
        "import frida,json;"
        "print(json.dumps([{'id':d.id,'name':d.name,'type':d.type} "
        "for d in frida.enumerate_devices() if d.type=='usb']))"
    )
    completed = subprocess.run(
        [frida_python(), "-c", code],
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
        env=child_env(),
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or "Không liệt kê được Frida devices.")
    return json.loads(completed.stdout or "[]")


def resolve_device_ids(value: str) -> list[str]:
    raw = (value or "auto").strip()
    lowered = raw.lower()
    if lowered in ("all", "*", "tat-ca", "tất-cả"):
        devices = list_usb_devices()
        ids = [str(device.get("id") or "").strip() for device in devices if device.get("id")]
        if not ids:
            raise RuntimeError("Không thấy iPhone USB nào qua Frida.")
        return ids
    if lowered == "auto":
        devices = list_usb_devices()
        ids = [str(device.get("id") or "").strip() for device in devices if device.get("id")]
        if not ids:
            raise RuntimeError("Không thấy iPhone USB nào qua Frida.")
        return ids
    ids = [part.strip() for part in raw.split(",") if part.strip()]
    return ids or ["auto"]


def device_label(device_id: str) -> str:
    if device_id == "auto":
        return "auto"
    return device_id[:8]


def safe_name(value: str) -> str:
    text = "".join(char if char.isalnum() or char in "._-" else "_" for char in str(value or ""))
    return text[:64] or "worker"


def ensure_result_headers(ws) -> dict[str, int]:
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    for header in ROW_RESULT_COLUMNS:
        if header not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=header)
            headers.append(header)
    return {header: index + 1 for index, header in enumerate(headers) if header}


def atomic_save_workbook(wb, xlsx: Path) -> None:
    tmp = xlsx.with_name(f".{xlsx.stem}.chiumerge.{os.getpid()}.{time.time_ns()}.tmp.xlsx")
    try:
        wb.save(tmp)
        tmp.replace(xlsx)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass


def collect_run_events(run_dir: Path) -> dict[tuple[str, int], dict]:
    records: dict[tuple[str, int], dict] = {}
    for path in sorted(run_dir.glob("*.jsonl")):
        with path.open("r", encoding="utf-8", errors="replace") as handle:
            for line in handle:
                line = line.strip()
                if not line:
                    continue
                try:
                    event = json.loads(line)
                except json.JSONDecodeError:
                    continue
                sheet = str(event.get("sheet") or "")
                try:
                    row = int(event.get("row") or 0)
                except (TypeError, ValueError):
                    continue
                if not sheet or row <= 0:
                    continue
                key = (sheet, row)
                current = records.setdefault(
                    key,
                    {
                        "sheet": sheet,
                        "row": row,
                        "device_id": str(event.get("device_id") or ""),
                        "crane_result": {},
                    },
                )
                stage = str(event.get("stage") or "")
                if event.get("device_id"):
                    current["device_id"] = str(event.get("device_id") or "")
                if isinstance(event.get("crane_result"), dict):
                    current["crane_result"] = event["crane_result"]
                if stage == "final":
                    current["status"] = str(event.get("status") or "FAILED").strip().upper()
                    current["error_details"] = str(event.get("error_details") or "")
                    current["reg_status"] = str(event.get("reg_status") or "")
                    current["gold_status"] = str(event.get("gold_status") or "")
                    current["chiu_status"] = str(event.get("chiu_status") or "")
                    current["notes"] = str(event.get("notes") or "")
                    current["final_seen"] = True
    return records


def merge_run_events_to_excel(xlsx: Path, run_dir: Path, default_sheet: str) -> tuple[int, int]:
    records = collect_run_events(run_dir)
    if not records:
        return 0, 0
    merged = 0
    partial = 0
    with excel_write_lock(xlsx):
        lock_path = xlsx.parent / f".~lock.{xlsx.name}#"
        if lock_path.exists():
            raise RuntimeError(f"Excel appears to be open/locked: {lock_path}. Đóng file rồi chạy lại.")
        wb = openpyxl.load_workbook(xlsx)
        try:
            for (sheet_name, row), record in sorted(records.items(), key=lambda item: (item[0][0], item[0][1])):
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                elif default_sheet in wb.sheetnames:
                    ws = wb[default_sheet]
                else:
                    ws = choose_sheet(wb, default_sheet)
                col = ensure_result_headers(ws)
                crane = record.get("crane_result") or {}
                now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                values = {
                    "crane_container_id": crane.get("crane_container_id") or crane.get("active_container_id") or "",
                    "crane_container_name": crane.get("crane_container_name") or crane.get("active_container_name") or "",
                    "crane_status": "ASSIGNED" if crane else "",
                    "crane_assigned_at": now if crane else "",
                    "crane_last_used_at": now if crane else "",
                    "frida_device_id": crane.get("frida_device_id") or record.get("device_id") or "",
                    "frida_device_name": crane.get("frida_device_name") or "",
                }
                if record.get("final_seen"):
                    values["status"] = record.get("status") or "FAILED"
                    values["error_details"] = record.get("error_details") or ""
                    if record.get("reg_status"):
                        values["reg_status"] = record.get("reg_status")
                    if record.get("gold_status"):
                        values["gold_status"] = record.get("gold_status")
                    if record.get("chiu_status"):
                        values["chiu_status"] = record.get("chiu_status")
                    if record.get("notes"):
                        values["notes"] = record.get("notes")
                    merged += 1
                else:
                    partial += 1
                for key, value in values.items():
                    if value in (None, "") and key not in ("error_details",):
                        continue
                    if key == "crane_assigned_at" and ws.cell(row=row, column=col[key]).value:
                        continue
                    ws.cell(row=row, column=col[key], value=value)
            atomic_save_workbook(wb, xlsx)
        finally:
            wb.close()
    return merged, partial


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run Yamada_chiu flow for multiple rows in one sheet.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Accounts")
    parser.add_argument("--limit", type=int, default=0, help="0 = run all runnable rows; N = run first N runnable rows.")
    parser.add_argument("--wait-timeout-ms", type=int, default=25000)
    parser.add_argument("--max-steps", type=int, default=40)
    parser.add_argument("--flow-complete-delay-ms", type=int, default=2000)
    parser.add_argument("--gold-payment-wait-timeout-ms", type=int, default=60000)
    parser.add_argument("--device-id", default=os.environ.get("FRIDA_DEVICE_ID", "auto"), help="'auto', 'all', or comma-separated Frida device IDs.")
    parser.add_argument("--no-reload", action="store_true")
    parser.add_argument("--no-submit", action="store_true")
    parser.add_argument(
        "--new-container",
        action="store_true",
        help="Tạo container Crane MỚI cho mỗi nick (đăng ký/đăng nhập) và ghi containerID + deviceID ra Excel. "
        "Không bật thì giữ flow cũ (reuse container active, chỉ đăng nhập nếu cần).",
    )
    parser.add_argument("--max-attempts", type=int, default=2, help="Max attempts per row, including the first run.")
    parser.add_argument(
        "--respring-every",
        type=int,
        default=5,
        help="Clean-respring each phone after every N rows it ran (frees RAM before jetsam storms panic launchd). 0 = off.",
    )
    parser.add_argument("--row-range", default="", help="Only run rows A-B (inclusive), e.g. 640-969.")
    parser.add_argument(
        "--create-container",
        action="store_true",
        help="Always create a new Crane container per row (login flow, no register). For moving rows to another phone.",
    )
    parser.add_argument("--list-only", action="store_true")
    parser.add_argument("--direct-excel-write", action="store_true", help="Old mode: each worker writes Excel after every row.")
    parser.add_argument("--merge-run-dir", default="", help="Merge a previous batch run_dir JSONL into Excel, then exit.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    xlsx = Path(args.xlsx).expanduser()
    if not xlsx.exists():
        print(f"[batch] Không thấy file Excel: {xlsx}", file=sys.stderr, flush=True)
        return 1
    if args.limit < 0:
        print("[batch] Số nick phải >= 0.", file=sys.stderr, flush=True)
        return 1
    if args.max_attempts < 1:
        print("[batch] max-attempts phải >= 1.", file=sys.stderr, flush=True)
        return 1
    if args.merge_run_dir:
        try:
            merged, partial = merge_run_events_to_excel(xlsx, Path(args.merge_run_dir).expanduser(), args.sheet)
            print(f"[batch] Đã merge {merged} final row vào Excel. Partial/no-final={partial}.", flush=True)
            return 0
        except Exception as exc:
            print(f"[batch] Merge run_dir lỗi: {exc}", file=sys.stderr, flush=True)
            return 1
    try:
        device_ids = resolve_device_ids(args.device_id)
    except Exception as exc:
        print(f"[batch] {exc}", file=sys.stderr, flush=True)
        return 1

    requested_device = (args.device_id or "auto").strip().lower()
    allowed_device_ids = set(device_ids)
    balance_device_ids = device_ids if len(device_ids) > 1 else None
    sheet, tasks = runnable_rows(xlsx, args.sheet, args.limit, allowed_device_ids, balance_device_ids)
    if args.row_range:
        lo, hi = (int(x) for x in args.row_range.split("-"))
        tasks = [task for task in tasks if lo <= int(task["row"]) <= hi]
    print(f"[batch] Sheet={sheet} | số nick={'full' if args.limit == 0 else args.limit} | chọn {len(tasks)} row", flush=True)
    print(f"[batch] Devices: {', '.join(device_ids)}", flush=True)
    if allowed_device_ids is not None:
        print(f"[batch] Lọc theo frida_device_id trong Excel: {', '.join(sorted(allowed_device_ids))}", flush=True)
    if tasks:
        preview = ", ".join(str(task["row"]) for task in tasks[:20])
        suffix = "..." if len(tasks) > 20 else ""
        print(f"[batch] Rows: {preview}{suffix}", flush=True)
    if args.list_only:
        return 0
    if not tasks:
        print("[batch] Không có row PENDING/FAILED/trống có email để chạy.", flush=True)
        return 0

    failures: list[tuple[int, int]] = []
    success_durations: list[float] = []
    batch_start = time.monotonic()
    print_lock = threading.Lock()
    result_lock = threading.Lock()
    stop_all = threading.Event()
    total_tasks = len(tasks)
    run_dir = ROOT_DIR / "agents" / "runtime" / "batch_runs" / f"{datetime.now():%Y%m%d_%H%M%S}_{os.getpid()}"
    if not args.direct_excel_write:
        run_dir.mkdir(parents=True, exist_ok=True)
        print(f"[batch] Deferred Excel write: worker logs ở {run_dir}", flush=True)
    card_device_ids = ",".join(device_ids)
    task_queues: dict[str, queue.Queue[dict]] = {device_id: queue.Queue() for device_id in device_ids}
    queued_per_device = {device_id: 0 for device_id in device_ids}
    rr = 0
    for task in tasks:
        preferred = str(task.get("device_id") or "")
        if preferred in task_queues:
            target = preferred
        else:
            target = device_ids[rr % len(device_ids)]
            rr += 1
        queued_task = dict(task)
        if args.new_container:
            # Checkbox "Tạo container mới mỗi nick": luôn tạo container Crane mới,
            # bỏ qua container cũ trên row; crane ghi lại containerID + deviceID ra Excel.
            queued_task["container_mode"] = "create"
        elif args.create_container:
            queued_task["container_mode"] = "create"
        else:
            queued_task["container_mode"] = (
                "active-then-create"
                if task.get("container_id") or queued_per_device[target] == 0
                else "create"
            )
        queued_per_device[target] += 1
        task_queues[target].put(queued_task)

    def log(line: str = "") -> None:
        with print_lock:
            print(line, flush=True)

    def run_task(device_id: str, index: int, row: int, container_mode: str) -> tuple[int, str, float]:
        cmd = [
            sys.executable,
            "scripts/chiu_full_flow.py",
            "--xlsx",
            str(xlsx),
            "--sheet",
            sheet,
            "--row",
            str(row),
            "--wait-timeout-ms",
            str(args.wait_timeout_ms),
            "--max-steps",
            str(args.max_steps),
            "--flow-complete-delay-ms",
            str(args.flow_complete_delay_ms),
            "--gold-payment-wait-timeout-ms",
            str(args.gold_payment_wait_timeout_ms),
            "--device-id",
            device_id,
            "--container-mode",
            container_mode,
            "--card-device-ids",
            card_device_ids,
        ]
        if not args.direct_excel_write:
            cmd.extend([
                "--defer-excel-write",
                "--event-log",
                str(run_dir / f"{safe_name(device_id)}.jsonl"),
            ])
        if args.no_reload:
            cmd.append("--no-reload")
        if args.no_submit:
            cmd.append("--no-submit")
        if args.new_container:
            # Checkbox "luồng mới": bật luồng đăng ký (reg001->reg006) trong DOM.
            # container_mode="create" (ở trên) đã lo phần tạo container mới.
            cmd.append("--enable-register")

        final_code = 0
        final_output = ""
        row_start = time.monotonic()
        for attempt in range(1, args.max_attempts + 1):
            log(
                f"\n[batch][{device_label(device_id)}] ({index}/{total_tasks}) chạy row {row} "
                f"| container-mode={container_mode} | attempt {attempt}/{args.max_attempts}"
            )
            log(f"[batch][{device_label(device_id)}] $ " + quote_cmd(cmd))
            proc = subprocess.Popen(
                cmd,
                cwd=str(ROOT_DIR),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                bufsize=1,
                env=child_env(),
            )
            assert proc.stdout is not None
            output_lines: list[str] = []
            for line in proc.stdout:
                output_lines.append(line)
                log(f"[{device_label(device_id)} r{row}] {line.rstrip()}")
            final_code = proc.wait()
            final_output = "".join(output_lines)
            if final_code == 0:
                break
            if attempt < args.max_attempts:
                log(f"[batch][{device_label(device_id)}] Row {row} lỗi exit={final_code}, retry lần cuối...")
            if stop_all.is_set():
                break

        row_elapsed = time.monotonic() - row_start
        return final_code, final_output, row_elapsed

    counter_lock = threading.Lock()
    counter = {"value": 0}

    def respring(device_id: str) -> None:
        cmd = [sys.executable, "scripts/phone_ram_guard.py", "--device-id", device_id, "--force-respring"]
        log(f"\n[batch][{device_label(device_id)}] Respring định kỳ (mỗi {args.respring_every} nick)")
        try:
            done = subprocess.run(
                cmd, cwd=str(ROOT_DIR), capture_output=True, text=True,
                encoding="utf-8", errors="replace", env=child_env(), timeout=180,
            )
            log(f"[batch][{device_label(device_id)}] {(done.stdout + done.stderr).strip().splitlines()[-1:]}")
        except Exception as exc:
            log(f"[batch][{device_label(device_id)}] Respring lỗi: {exc}")

    def worker(device_id: str) -> None:
        rows_done = 0
        while not stop_all.is_set():
            try:
                task = task_queues[device_id].get_nowait()
            except queue.Empty:
                return
            if args.respring_every > 0 and rows_done and rows_done % args.respring_every == 0:
                respring(device_id)
            rows_done += 1
            with counter_lock:
                counter["value"] += 1
                index = counter["value"]
            row = int(task["row"])
            container_mode = str(task.get("container_mode") or "active-then-next")
            final_code, final_output, row_elapsed = run_task(device_id, index, row, container_mode)
            label = device_label(device_id)
            with result_lock:
                if final_code != 0:
                    failures.append((row, final_code))
                    if any(hint in final_output for hint in STOP_ON_ERROR_HINTS):
                        # Only this device stops; other devices keep running.
                        log(f"[batch][{label}] Dừng máy này vì lỗi hạ tầng ở row {row} sau {args.max_attempts} attempt.")
                        return
                    else:
                        log(f"[batch][{label}] Row {row} lỗi exit={final_code} sau {args.max_attempts} attempt, chuyển row tiếp theo.")
                else:
                    success_durations.append(row_elapsed)
                    avg = sum(success_durations) / len(success_durations)
                    log(f"[batch][{label}] Row {row} xong trong {row_elapsed:.1f}s | trung bình {avg:.1f}s/nick")
            if not stop_all.is_set() and not task_queues[device_id].empty():
                time.sleep(random.uniform(2.0, 3.0))

    threads = [threading.Thread(target=worker, args=(device_id,), daemon=True) for device_id in device_ids]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    if not args.direct_excel_write:
        try:
            merged, partial = merge_run_events_to_excel(xlsx, run_dir, sheet)
            print(f"\n[batch] Đã ghi Excel cuối batch: {merged} row final | partial/no-final={partial}", flush=True)
            print(f"[batch] Run log giữ tại: {run_dir}", flush=True)
        except Exception as exc:
            print(f"\n[batch] Lỗi merge Excel cuối batch: {exc}", file=sys.stderr, flush=True)
            print(f"[batch] Data tạm vẫn còn ở: {run_dir}", file=sys.stderr, flush=True)
            return 1

    if failures:
        detail = ", ".join(f"row {row}: exit {code}" for row, code in failures[:10])
        more = "..." if len(failures) > 10 else ""
        print(f"\n[batch] Xong, có {len(failures)} row lỗi: {detail}{more}", flush=True)
    else:
        print("\n[batch] Xong, không có row lỗi.", flush=True)
    if success_durations:
        total = time.monotonic() - batch_start
        avg = sum(success_durations) / len(success_durations)
        per_hour = 3600 / avg if avg > 0 else 0
        print(f"[batch] Thống kê: {len(success_durations)} nick OK | avg={avg:.1f}s/nick | ~{per_hour:.1f} nick/giờ | total={total:.1f}s", flush=True)
    return 0


if __name__ == "__main__":
    configure_stdio()
    raise SystemExit(main())
