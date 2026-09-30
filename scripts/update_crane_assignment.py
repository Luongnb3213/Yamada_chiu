from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import COLUMN_ALIASES, INPUT_HEADERS  # noqa: E402


CRANE_COLUMNS = [
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "crane_assigned_at",
    "crane_last_used_at",
]


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    return COLUMN_ALIASES.get(text, text)


def load_jsonish(value: str) -> dict:
    if not value:
        return {}
    path = Path(value).expanduser()
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return json.loads(value)


def ensure_headers(ws) -> dict[str, int]:
    headers = [normalize_header(cell.value) for cell in ws[1]]
    if not headers or "email" not in headers:
        ws.delete_rows(1, ws.max_row)
        ws.append(INPUT_HEADERS)
        headers = [normalize_header(cell.value) for cell in ws[1]]

    for col in CRANE_COLUMNS:
        if col not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=col)
            headers.append(col)
    return {header: idx + 1 for idx, header in enumerate(headers) if header}


def find_row(ws, col_map: dict[str, int], row_number: int | None, email: str) -> int:
    if row_number:
        return row_number
    wanted = email.strip().lower()
    if not wanted:
        raise RuntimeError("Need --row or --email.")
    email_col = col_map.get("email")
    if not email_col:
        raise RuntimeError("Sheet does not have an email column.")
    for idx in range(2, ws.max_row + 1):
        current = str(ws.cell(row=idx, column=email_col).value or "").strip().lower()
        if current == wanted:
            return idx
    raise RuntimeError(f"Email not found: {email}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Write Crane container assignment back to a Yamada Excel row.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Accounts")
    parser.add_argument("--row", type=int)
    parser.add_argument("--email", default="")
    parser.add_argument("--container-id", default="")
    parser.add_argument("--container-name", default="")
    parser.add_argument("--status", default="ASSIGNED")
    parser.add_argument("--from-json", default="", help="JSON string/file returned by craneEnsureForProfile().")
    args = parser.parse_args()

    payload = load_jsonish(args.from_json) if args.from_json else {}
    container_id = args.container_id or payload.get("crane_container_id") or payload.get("active_container_id") or ""
    container_name = args.container_name or payload.get("crane_container_name") or ""
    if not container_id:
        raise SystemExit("Missing container id. Pass --container-id or --from-json.")

    xlsx = Path(args.xlsx).expanduser()
    if not xlsx.exists():
        raise SystemExit(f"XLSX not found: {xlsx}")

    wb = openpyxl.load_workbook(xlsx)
    try:
        if args.sheet in wb.sheetnames:
            ws = wb[args.sheet]
        elif "Accounts" in wb.sheetnames:
            ws = wb["Accounts"]
        elif "Mail" in wb.sheetnames:
            ws = wb["Mail"]
        else:
            ws = wb[wb.sheetnames[0]]
        col_map = ensure_headers(ws)
        target_row = find_row(ws, col_map, args.row, args.email)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        ws.cell(row=target_row, column=col_map["crane_container_id"], value=container_id)
        ws.cell(row=target_row, column=col_map["crane_container_name"], value=container_name)
        ws.cell(row=target_row, column=col_map["crane_status"], value=args.status)
        if not ws.cell(row=target_row, column=col_map["crane_assigned_at"]).value:
            ws.cell(row=target_row, column=col_map["crane_assigned_at"], value=now)
        ws.cell(row=target_row, column=col_map["crane_last_used_at"], value=now)

        tmp = str(xlsx) + ".tmp"
        wb.save(tmp)
        Path(tmp).replace(xlsx)
    finally:
        wb.close()

    print(json.dumps({
        "ok": True,
        "xlsx": str(xlsx),
        "sheet": ws.title,
        "row": target_row,
        "crane_container_id": container_id,
        "crane_container_name": container_name,
        "crane_status": args.status,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
