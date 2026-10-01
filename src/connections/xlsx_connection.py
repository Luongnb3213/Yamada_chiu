from __future__ import annotations

import os
import threading
import time
import zipfile
import zlib
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import config
from src.utils.logger import get_logger


log = get_logger("xlsx_connection")
_FALLBACK_XLSX_LOCK = threading.Lock()


INPUT_HEADERS = [
    "email",
    "password",
    "otp_email",
    "otp_pass",
    "otp_imap_host",
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "crane_assigned_at",
    "crane_last_used_at",
    "frida_device_id",
    "frida_device_name",
    "gold_status",
    "chiu_status",
    "status",
    "error_details",
    "notes",
    "credit_card_number",
    "credit_card_exp",
    "credit_card_cvv",
    "onepiece_shop_name",
    "onepiece_area",
]
RESULT_HEADERS = [
    "email",
    "password",
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "frida_device_id",
    "frida_device_name",
    "proxy_used",
    "proxy_id",
    "status",
    "created_at",
    "error_details",
]
PROXIES_HEADERS = ["proxy", "status", "proxy_id"]

COLUMN_ALIASES = {
    "メールアドレス": "email",
    "email": "email",
    "email address": "email",
    "mail": "email",
    "account": "email",
    "account_email": "email",
    "pass": "password",
    "email_password": "password",
    "email password": "password",
    "mail_password": "password",
    "mail password": "password",
    "account_password": "password",
    "otp_email": "otp_email",
    "otp email": "otp_email",
    "otp_mail": "otp_email",
    "otp_pass": "otp_pass",
    "otp pass": "otp_pass",
    "otp_password": "otp_pass",
    "otp_imap_host": "otp_imap_host",
    "imap_host": "otp_imap_host",
    "proxy used": "proxy_used",
    "proxy id": "proxy_id",
    "crane container id": "crane_container_id",
    "crane_container_id": "crane_container_id",
    "container_id": "crane_container_id",
    "crane id": "crane_container_id",
    "crane_id": "crane_container_id",
    "crane container name": "crane_container_name",
    "crane_container_name": "crane_container_name",
    "container_name": "crane_container_name",
    "crane name": "crane_container_name",
    "crane_name": "crane_container_name",
    "crane status": "crane_status",
    "crane_status": "crane_status",
    "crane assigned at": "crane_assigned_at",
    "crane_assigned_at": "crane_assigned_at",
    "crane last used at": "crane_last_used_at",
    "crane_last_used_at": "crane_last_used_at",
    "frida device id": "frida_device_id",
    "frida_device_id": "frida_device_id",
    "device_id": "frida_device_id",
    "frida device name": "frida_device_name",
    "frida_device_name": "frida_device_name",
    "device_name": "frida_device_name",
    "gold_status": "gold_status",
    "gold status": "gold_status",
    "membership_status": "gold_status",
    "gold_membership_status": "gold_status",
    "chiu_status": "chiu_status",
    "chiu status": "chiu_status",
    "credit_card_number": "credit_card_number",
    "card_number": "credit_card_number",
    "card number": "credit_card_number",
    "credit_card_exp": "credit_card_exp",
    "card_exp": "credit_card_exp",
    "card exp": "credit_card_exp",
    "credit_card_cvv": "credit_card_cvv",
    "card_cvv": "credit_card_cvv",
    "cvv": "credit_card_cvv",
    "onepiece_shop_name": "onepiece_shop_name",
    "onepiece shop name": "onepiece_shop_name",
    "lottery_shop_name": "onepiece_shop_name",
    "shop_name": "onepiece_shop_name",
    "onepiece_area": "onepiece_area",
    "onepiece area": "onepiece_area",
    "lottery_area": "onepiece_area",
    "created at": "created_at",
    "error details": "error_details",
    "reg_status": "reg_status",
    "reg status": "reg_status",
    "register_status": "reg_status",
    "registration_status": "reg_status",
    # Registration (reg001 -> reg006) member-info fields.
    "pin": "pin",
    "pass_pin": "pin",
    "login_pin": "pin",
    "phone": "phone",
    "tel": "phone",
    "phone_number": "phone",
    "last_name": "last_name",
    "last name": "last_name",
    "sei": "last_name",
    "first_name": "first_name",
    "first name": "first_name",
    "mei": "first_name",
    "last_name_kana": "last_name_kana",
    "last name kana": "last_name_kana",
    "seikana": "last_name_kana",
    "first_name_kana": "first_name_kana",
    "first name kana": "first_name_kana",
    "meikana": "first_name_kana",
    "postal_code": "postal_code",
    "postal code": "postal_code",
    "zip": "postal_code",
    "zipcode": "postal_code",
    "prefecture": "prefecture",
    "prefcode": "prefecture",
    "city": "city",
    "adrs1": "city",
    "address_rest": "address_rest",
    "address": "address_rest",
    "adrs2": "address_rest",
    "dob": "dob",
    "birthday": "dob",
    "birth_date": "dob",
    "birthdate": "dob",
    "gender": "gender",
    "sex": "gender",
}

STATUS_MAP = {
    "": "PENDING",
    "pending": "PENDING",
    "processing": "PROCESSING",
    "success": "SUCCESS",
    "failed": "FAILED",
    "error": "FAILED",
    "fail_no_retry": "FAIL_NO_RETRY",
    "aborted": "FAIL_NO_RETRY",
}


def normalize_status(status: str) -> str:
    return STATUS_MAP.get(str(status or "").strip().lower(), str(status or "FAILED").strip().upper())


@contextmanager
def excel_write_lock(xlsx_path: str | Path):
    """Serialize writes to one XLSX across parallel device workers."""
    path = Path(xlsx_path).expanduser()
    lock_path = path.parent / f".{path.name}.write.lock"
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    handle = lock_path.open("a+")
    locked_by = ""
    try:
        if os.name == "nt":
            import msvcrt

            deadline = time.monotonic() + 300
            while True:
                try:
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                    locked_by = "msvcrt"
                    break
                except OSError:
                    if time.monotonic() >= deadline:
                        raise TimeoutError(f"Timeout waiting for Excel write lock: {lock_path}")
                    time.sleep(0.2)
        else:
            try:
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
                locked_by = "fcntl"
            except Exception:
                _FALLBACK_XLSX_LOCK.acquire()
                locked_by = "thread"
        yield
    finally:
        try:
            if locked_by == "msvcrt":
                import msvcrt

                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            elif locked_by == "fcntl":
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            elif locked_by == "thread":
                _FALLBACK_XLSX_LOCK.release()
        except Exception:
            if locked_by == "thread":
                try:
                    _FALLBACK_XLSX_LOCK.release()
                except RuntimeError:
                    pass
        finally:
            handle.close()


def repair_xlsx_crc(corrupted_file_path: str, output_file_path: str) -> None:
    """Rebuild an XLSX zip archive when Excel left a bad CRC behind."""
    with zipfile.ZipFile(corrupted_file_path, "r") as zin:
        with zipfile.ZipFile(output_file_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = None
                try:
                    fp = zin.open(item)
                    if hasattr(fp, "_expected_crc"):
                        fp._expected_crc = None
                    content = fp.read()
                except Exception:
                    try:
                        zinfo = zin.getinfo(item.filename)
                        with open(corrupted_file_path, "rb") as f:
                            f.seek(zinfo.header_offset)
                            header = f.read(30)
                            fname_len = int.from_bytes(header[26:28], "little")
                            extra_len = int.from_bytes(header[28:30], "little")
                            f.seek(zinfo.header_offset + 30 + fname_len + extra_len)
                            compressed_data = f.read(zinfo.compress_size)
                            content = (
                                zlib.decompress(compressed_data, -15)
                                if zinfo.compress_type == zipfile.ZIP_DEFLATED
                                else compressed_data
                            )
                    except Exception as exc:
                        log.warning("Không thể đọc file %s trong archive: %s", item.filename, exc)
                        continue
                if content is not None:
                    new_info = zipfile.ZipInfo(item.filename, item.date_time)
                    new_info.compress_type = zipfile.ZIP_DEFLATED
                    zout.writestr(new_info, content)


def cleanup_empty_rows(ws) -> int:
    if ws.max_row < 2:
        return 0
    empty_rows = []
    max_c = max(ws.max_column, 10)
    for row_idx in range(2, ws.max_row + 1):
        if all(
            ws.cell(row=row_idx, column=col_idx).value in (None, "")
            or str(ws.cell(row=row_idx, column=col_idx).value).strip() == ""
            for col_idx in range(1, max_c + 1)
        ):
            empty_rows.append(row_idx)
    for row_idx in reversed(empty_rows):
        ws.delete_rows(row_idx)
    if empty_rows:
        log.info("Đã dọn %s dòng trống trong sheet %s.", len(empty_rows), ws.title)
    return len(empty_rows)


class XlsxConnection:
    """Thread-safe local XLSX connection for Yamada_chiu wrapper code."""

    def __init__(self, xlsx_path: str):
        self.xlsx_path = Path(xlsx_path) if xlsx_path else None
        self._lock = threading.Lock()
        self._connected = False

        if not self.xlsx_path:
            log.warning("Chưa cấu hình đường dẫn file XLSX.")
            return
        if not self.xlsx_path.exists():
            log.warning("File XLSX không tồn tại: %s", self.xlsx_path)
            return

        try:
            self._create_session_backup()
            wb = self._load_workbook()
            self._ensure_sheets(wb)
            self._atomic_save(wb, self.xlsx_path)
            wb.close()
            self._connected = True
            log.info("XlsxConnection kết nối thành công: %s", self.xlsx_path)
        except Exception as exc:
            log.error("Không thể mở file XLSX: %s", exc)

    def is_connected(self) -> bool:
        return self._connected

    def _atomic_save(self, wb, path: Path) -> None:
        tmp_path = str(path) + f".tmp_{int(datetime.now().timestamp() * 1000)}"
        try:
            wb.save(tmp_path)
            test_wb = openpyxl.load_workbook(tmp_path, read_only=True)
            test_wb.close()
            os.replace(tmp_path, str(path))
        except Exception as exc:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass
            try:
                wb.save(str(path))
            except Exception as direct_exc:
                raise RuntimeError(
                    f"Không thể lưu file Excel. Hãy đóng file nếu đang mở: {direct_exc}"
                ) from exc

    def _create_session_backup(self) -> None:
        if not self.xlsx_path or not self.xlsx_path.exists():
            return
        try:
            backup_dir = self.xlsx_path.parent / "backups"
            backup_dir.mkdir(exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_file = backup_dir / f"{self.xlsx_path.stem}_backup_{timestamp}{self.xlsx_path.suffix}"
            import shutil

            shutil.copy2(self.xlsx_path, backup_file)
            backups = sorted(backup_dir.glob(f"{self.xlsx_path.stem}_backup_*"), key=os.path.getmtime)
            for old_backup in backups[:-10]:
                try:
                    old_backup.unlink()
                except Exception:
                    pass
            log.info("Đã tạo backup XLSX: %s", backup_file)
        except Exception as exc:
            log.warning("Không thể tạo backup XLSX: %s", exc)

    def _load_workbook(self, read_only: bool = False):
        if not self.xlsx_path:
            raise RuntimeError("Chưa cấu hình đường dẫn file XLSX.")
        try:
            return openpyxl.load_workbook(str(self.xlsx_path), read_only=read_only)
        except Exception as exc:
            err = str(exc)
            if any(token in err for token in ("Bad CRC-32", "BadZipFile", "not a zip file", "zipfile")):
                log.error("Phát hiện XLSX hỏng, thử khôi phục: %s", self.xlsx_path)
                repaired_tmp = str(self.xlsx_path) + ".repaired"
                repair_xlsx_crc(str(self.xlsx_path), repaired_tmp)
                os.replace(repaired_tmp, str(self.xlsx_path))
                return openpyxl.load_workbook(str(self.xlsx_path), read_only=read_only)
            raise

    @staticmethod
    def _normalize_header(value: object) -> str:
        text = str(value or "").strip().lower()
        return COLUMN_ALIASES.get(text, text)

    @classmethod
    def _get_headers(cls, ws) -> list[str]:
        if ws.max_row < 1:
            return []
        return [cls._normalize_header(cell.value) for cell in ws[1]]

    @staticmethod
    def _col_index(headers: list[str], col_name: str) -> int:
        target = COLUMN_ALIASES.get(col_name.lower(), col_name.lower())
        if target not in headers:
            raise KeyError(f"Không tìm thấy cột '{col_name}'. Headers: {headers}")
        return headers.index(target)

    @staticmethod
    def _write_header(ws, headers: list[str], color: str = "4472C4") -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center")

    @classmethod
    def _ensure_header_columns(cls, ws, required_headers: list[str]) -> bool:
        headers = cls._get_headers(ws)
        if not headers:
            cls._write_header(ws, required_headers)
            return True

        changed = False
        for header in required_headers:
            normalized = cls._normalize_header(header)
            if normalized not in headers:
                ws.cell(row=1, column=len(headers) + 1, value=header)
                headers.append(normalized)
                changed = True
        return changed

    def _ensure_sheets(self, wb) -> None:
        if "Accounts" not in wb.sheetnames:
            ws = wb.create_sheet("Accounts")
            self._write_header(ws, INPUT_HEADERS, "4472C4")
        else:
            self._ensure_header_columns(wb["Accounts"], INPUT_HEADERS)
        if "Results" not in wb.sheetnames:
            ws = wb.create_sheet("Results")
            self._write_header(ws, RESULT_HEADERS, "70AD47")
        else:
            self._ensure_header_columns(wb["Results"], RESULT_HEADERS)
        if "Proxies" not in wb.sheetnames:
            ws = wb.create_sheet("Proxies")
            self._write_header(ws, PROXIES_HEADERS, "ED7D31")
        else:
            self._ensure_header_columns(wb["Proxies"], PROXIES_HEADERS)

        for sheet_name in wb.sheetnames:
            cleanup_empty_rows(wb[sheet_name])

    def _active_sheet_name(self, wb) -> str:
        active = str(getattr(config, "ACTIVE_SHEET", "Accounts") or "Accounts")
        if active in wb.sheetnames:
            return active
        if "Accounts" in wb.sheetnames:
            return "Accounts"
        return "Mail" if "Mail" in wb.sheetnames else wb.sheetnames[0]

    def reset_processing_to_pending(self) -> None:
        with self._lock:
            try:
                wb = self._load_workbook()
                ws = wb[self._active_sheet_name(wb)]
                headers = self._get_headers(ws)
                status_col = self._col_index(headers, "status")
                reset_count = 0
                for row in ws.iter_rows(min_row=2):
                    if normalize_status(row[status_col].value) == "PROCESSING":
                        row[status_col].value = "PENDING"
                        reset_count += 1
                if reset_count:
                    self._atomic_save(wb, self.xlsx_path)
                    log.info("Reset %s dòng PROCESSING về PENDING.", reset_count)
                wb.close()
            except Exception as exc:
                log.error("Lỗi reset PROCESSING: %s", exc)

    reset_interrupted_to_pending = reset_processing_to_pending

    def get_pending_accounts(self, batch_size: int = 50) -> list[dict]:
        with self._lock:
            try:
                wb = self._load_workbook()
                ws = wb[self._active_sheet_name(wb)]
                headers = self._get_headers(ws)
                results = []
                for row in ws.iter_rows(min_row=2):
                    if len(results) >= batch_size:
                        break
                    row_values = [cell.value for cell in row]
                    row_dict = dict(zip(headers, row_values))
                    email_raw = str(row_dict.get("email") or "").strip()
                    if not email_raw:
                        continue
                    status = normalize_status(row_dict.get("status", "PENDING"))
                    if status not in ("PENDING", "FAILED"):
                        continue
                    parts = email_raw.split("|")
                    email = parts[0].strip()
                    password = (
                        parts[1].strip()
                        if len(parts) > 1
                        else str(row_dict.get("password") or "").strip()
                    )
                    account = {
                        "email": email,
                        "raw_email": email_raw,
                        "password": password,
                        "phone": str(row_dict.get("phone") or "").strip(),
                        "status": status,
                        "error_details": str(row_dict.get("error_details") or "").strip(),
                        "notes": str(row_dict.get("notes") or "").strip(),
                    }
                    for key, value in row_dict.items():
                        if key and key not in account:
                            account[key] = "" if value is None else value
                    results.append(account)
                wb.close()
                log.info("Đọc được %s account đang chờ từ XLSX.", len(results))
                return results
            except Exception as exc:
                log.error("Lỗi đọc pending accounts: %s", exc)
                return []

    get_pending_emails = get_pending_accounts

    def update_account_status(
        self,
        email: str,
        status: str,
        error_details: str = "",
        extra_data: dict | None = None,
    ) -> bool:
        email = str(email or "").strip()
        if not email:
            return False
        with self._lock:
            try:
                wb = self._load_workbook()
                ws = wb[self._active_sheet_name(wb)]
                headers = self._get_headers(ws)
                email_col = self._col_index(headers, "email")

                if "status" in headers:
                    status_col = self._col_index(headers, "status")
                else:
                    status_col = len(headers)
                    ws.cell(row=1, column=status_col + 1, value="status")
                    headers.append("status")

                if "error_details" in headers:
                    error_col = self._col_index(headers, "error_details")
                else:
                    error_col = len(headers)
                    ws.cell(row=1, column=error_col + 1, value="error_details")
                    headers.append("error_details")

                extra_cols = {}
                for key in (extra_data or {}):
                    key_norm = self._normalize_header(key)
                    if not key_norm:
                        continue
                    if key_norm in headers:
                        extra_cols[key_norm] = headers.index(key_norm)
                    else:
                        extra_cols[key_norm] = len(headers)
                        ws.cell(row=1, column=len(headers) + 1, value=key_norm)
                        headers.append(key_norm)

                matched = False
                for row in ws.iter_rows(min_row=2):
                    cell_raw = str(row[email_col].value or "").strip()
                    cell_email = cell_raw.split("|")[0].strip()
                    if cell_raw == email or cell_email == email:
                        row[status_col].value = normalize_status(status)
                        row[error_col].value = str(error_details or "")
                        for key, col_idx in extra_cols.items():
                            value = (extra_data or {}).get(key)
                            if value not in (None, ""):
                                ws.cell(row=row[0].row, column=col_idx + 1, value=str(value))
                        matched = True
                        break

                if matched:
                    self._atomic_save(wb, self.xlsx_path)
                wb.close()
                return matched
            except Exception as exc:
                log.error("Lỗi cập nhật account status: %s", exc)
                return False

    update_email_status = update_account_status

    def append_result(self, data: dict) -> None:
        email_raw = str(data.get("email") or "").strip()
        if not email_raw:
            return
        email = email_raw.split("|")[0].strip()
        status = normalize_status(data.get("status", "FAILED"))
        if status == "PROCESSING":
            return

        payload = dict(data)
        payload["email"] = email
        payload["status"] = status
        if status == "SUCCESS" and not payload.get("created_at"):
            payload["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with self._lock:
            try:
                wb = self._load_workbook()
                ws = wb["Results"] if "Results" in wb.sheetnames else wb.create_sheet("Results")
                headers = self._get_headers(ws)
                if not headers or "email" not in headers:
                    ws.delete_rows(1, ws.max_row)
                    self._write_header(ws, RESULT_HEADERS, "70AD47")
                    headers = self._get_headers(ws)

                for col_name in RESULT_HEADERS:
                    if col_name not in headers:
                        ws.cell(row=1, column=len(headers) + 1, value=col_name)
                        headers.append(col_name)

                for key in payload:
                    key_norm = self._normalize_header(key)
                    if key_norm and key_norm not in headers:
                        ws.cell(row=1, column=len(headers) + 1, value=key_norm)
                        headers.append(key_norm)

                cleanup_empty_rows(ws)
                col_map = {header: idx + 1 for idx, header in enumerate(headers)}
                email_col = col_map["email"]
                target_row = None
                for row in ws.iter_rows(min_row=2):
                    cell_email = str(ws.cell(row=row[0].row, column=email_col).value or "").strip().split("|")[0].strip()
                    if cell_email.lower() == email.lower():
                        target_row = row[0].row
                        break
                if target_row is None:
                    target_row = ws.max_row + 1

                for key, value in payload.items():
                    key_norm = self._normalize_header(key)
                    if key_norm in col_map and (value not in (None, "") or key_norm in ("status", "error_details")):
                        ws.cell(row=target_row, column=col_map[key_norm], value="" if value is None else value)

                self._atomic_save(wb, self.xlsx_path)
                wb.close()
                log.info("Đã ghi kết quả %s -> %s.", email, status)
            except Exception as exc:
                log.error("Không thể ghi Results cho %s: %s", email, exc, exc_info=True)

    append_account = append_result

    def get_result_info(self, email: str) -> dict:
        email = str(email or "").strip()
        if not email:
            return {}
        with self._lock:
            try:
                wb = self._load_workbook(read_only=True)
                if "Results" not in wb.sheetnames:
                    wb.close()
                    return {}
                ws = wb["Results"]
                headers = self._get_headers(ws)
                email_col = self._col_index(headers, "email")
                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_email = str(row[email_col] or "").strip().split("|")[0].strip()
                    if row_email.lower() == email.lower():
                        result = {
                            header: "" if idx >= len(row) or row[idx] is None else row[idx]
                            for idx, header in enumerate(headers)
                        }
                        wb.close()
                        return result
                wb.close()
                return {}
            except Exception as exc:
                log.error("Lỗi đọc result info: %s", exc)
                return {}

    get_account_info = get_result_info

    def get_active_proxies(self) -> list[dict]:
        with self._lock:
            try:
                wb = self._load_workbook()
                if "Proxies" not in wb.sheetnames:
                    wb.close()
                    return []
                ws = wb["Proxies"]
                headers = self._get_headers(ws)
                proxy_col = self._col_index(headers, "proxy")
                status_col = self._col_index(headers, "status") if "status" in headers else None

                changed = False
                if "proxy_id" in headers:
                    proxy_id_col = self._col_index(headers, "proxy_id")
                else:
                    proxy_id_col = len(headers)
                    ws.cell(row=1, column=proxy_id_col + 1, value="proxy_id")
                    headers.append("proxy_id")
                    changed = True

                results = []
                used_ids = set()
                next_id = 1
                for row_number in range(2, ws.max_row + 1):
                    proxy = str(ws.cell(row=row_number, column=proxy_col + 1).value or "").strip()
                    if not proxy:
                        continue
                    proxy_id = str(ws.cell(row=row_number, column=proxy_id_col + 1).value or "").strip()
                    if not proxy_id or proxy_id in used_ids:
                        while f"PX-{next_id:06d}" in used_ids:
                            next_id += 1
                        proxy_id = f"PX-{next_id:06d}"
                        next_id += 1
                        ws.cell(row=row_number, column=proxy_id_col + 1, value=proxy_id)
                        changed = True
                    used_ids.add(proxy_id)
                    if status_col is not None:
                        status = str(ws.cell(row=row_number, column=status_col + 1).value or "").strip().lower()
                        if status in ("disabled", "inactive", "used", "dead", "0"):
                            continue
                    results.append({"proxy_id": proxy_id, "raw": proxy, "sheet_row": row_number})
                if changed:
                    self._atomic_save(wb, self.xlsx_path)
                wb.close()
                log.info("Đọc được %s proxy active từ XLSX.", len(results))
                return results
            except Exception as exc:
                log.error("Lỗi đọc proxies: %s", exc)
                return []

    def update_proxy_status(self, proxy_ref, status: str = "INACTIVE") -> bool:
        return self.update_proxy_statuses([proxy_ref], status) > 0

    def update_proxy_statuses(self, proxy_refs, status: str = "INACTIVE") -> int:
        target_ids = set()
        target_rows = set()
        legacy_raws = []
        for ref in proxy_refs or []:
            if isinstance(ref, dict):
                proxy_id = str(ref.get("proxy_id") or "").strip()
                raw = str(ref.get("raw") or "").strip()
                row_number = ref.get("sheet_row")
                if proxy_id:
                    target_ids.add(proxy_id)
                elif isinstance(row_number, int) and row_number >= 2:
                    target_rows.add(row_number)
                elif raw:
                    legacy_raws.append(raw)
            else:
                raw = str(ref or "").strip()
                if raw:
                    legacy_raws.append(raw)
        if not target_ids and not target_rows and not legacy_raws:
            return 0

        with self._lock:
            try:
                wb = self._load_workbook()
                if "Proxies" not in wb.sheetnames:
                    wb.close()
                    return 0
                ws = wb["Proxies"]
                headers = self._get_headers(ws)
                proxy_col = self._col_index(headers, "proxy")
                proxy_id_col = self._col_index(headers, "proxy_id") if "proxy_id" in headers else None
                if "status" in headers:
                    status_col = self._col_index(headers, "status")
                else:
                    status_col = len(headers)
                    ws.cell(row=1, column=status_col + 1, value="status")

                matched = 0
                consumed_legacy = set()
                for row_number in range(2, ws.max_row + 1):
                    current = str(ws.cell(row=row_number, column=proxy_col + 1).value or "").strip()
                    current_id = (
                        str(ws.cell(row=row_number, column=proxy_id_col + 1).value or "").strip()
                        if proxy_id_col is not None
                        else ""
                    )
                    by_id = bool(current_id and current_id in target_ids)
                    by_row = row_number in target_rows
                    by_legacy_raw = current in legacy_raws and current not in consumed_legacy
                    if by_id or by_row or by_legacy_raw:
                        ws.cell(row=row_number, column=status_col + 1, value=str(status or "INACTIVE").upper())
                        matched += 1
                        if by_legacy_raw:
                            consumed_legacy.add(current)
                if matched:
                    self._atomic_save(wb, self.xlsx_path)
                    log.info("Đã cập nhật %s proxy sang %s.", matched, status)
                wb.close()
                return matched
            except Exception as exc:
                log.error("Không thể cập nhật proxy status: %s", exc)
                return 0

    def load_permanent_counts(self, proxy_pool) -> None:
        with self._lock:
            try:
                wb = self._load_workbook(read_only=True)
                if "Results" not in wb.sheetnames:
                    wb.close()
                    return
                ws = wb["Results"]
                headers = self._get_headers(ws)
                proxy_col = self._col_index(headers, "proxy_used")
                proxy_id_col = self._col_index(headers, "proxy_id") if "proxy_id" in headers else None
                counts = {}
                for row in ws.iter_rows(min_row=2, values_only=True):
                    proxy = str(row[proxy_col] or "").strip()
                    proxy_id = str(row[proxy_id_col] or "").strip() if proxy_id_col is not None else ""
                    identity = proxy_id or proxy
                    if identity and proxy.lower() != "direct":
                        counts[(identity, proxy)] = counts.get((identity, proxy), 0) + 1
                wb.close()
                for (identity, proxy_str), count in counts.items():
                    if hasattr(proxy_pool, "set_permanent_count"):
                        proxy_pool.set_permanent_count(
                            {"proxy_id": identity if identity.startswith("PX-") else "", "raw": proxy_str},
                            count,
                        )
                log.info("Đã khóa %s proxy từng được gán trong Results.", len(counts))
            except Exception as exc:
                log.warning("Không thể load permanent proxy counts: %s", exc)

    @staticmethod
    def create_template(xlsx_path: str) -> bool:
        try:
            wb = Workbook()
            ws_inputs = wb.active
            ws_inputs.title = "Accounts"
            XlsxConnection._write_header(ws_inputs, INPUT_HEADERS, "4472C4")
            ws_inputs.append([
                "example@example.com",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "PENDING",
                "",
                "",
            ])
            input_widths = [36, 22, 40, 28, 18, 22, 22, 48, 18, 14, 48, 32]
            for idx, width in enumerate(input_widths, 1):
                ws_inputs.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

            ws_results = wb.create_sheet("Results")
            XlsxConnection._write_header(ws_results, RESULT_HEADERS, "70AD47")
            for idx, width in enumerate((36, 22, 18, 34, 28, 18, 36, 16, 14, 22, 60), 1):
                ws_results.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

            ws_proxies = wb.create_sheet("Proxies")
            XlsxConnection._write_header(ws_proxies, PROXIES_HEADERS, "ED7D31")
            ws_proxies.append(["host:port:user:pass", "ACTIVE", "PX-000001"])
            for col, width in zip(("A", "B", "C"), (44, 14, 16)):
                ws_proxies.column_dimensions[col].width = width

            path = Path(xlsx_path)
            path.parent.mkdir(parents=True, exist_ok=True)
            tmp = str(path) + ".tmp"
            wb.save(tmp)
            os.replace(tmp, str(path))
            wb.close()
            log.info("Đã tạo XLSX template: %s", path)
            return True
        except Exception as exc:
            log.error("Lỗi tạo XLSX template: %s", exc)
            return False
